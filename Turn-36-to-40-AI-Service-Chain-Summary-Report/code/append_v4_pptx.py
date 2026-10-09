#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
總結報告_v4_補充版.pptx — AI 服務鏈總結報告生成腳本
策略：讀取 總結報告_v2.pptx，完全不更動前 17 頁，於最後面以「補充篇」形式加入 5 頁新投影片。
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os, sys

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# ─── Color Palette ───
NAVY       = RGBColor(0x0B, 0x1D, 0x3A)
DARK_BLUE  = RGBColor(0x14, 0x2D, 0x5E)
MID_BLUE   = RGBColor(0x1E, 0x3A, 0x6E)
ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xC1)
LIGHT_BLUE = RGBColor(0x5D, 0xAE, 0xF5)
GOLD       = RGBColor(0xF4, 0xC7, 0x30)
ORANGE     = RGBColor(0xE8, 0x7C, 0x2A)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xD5, 0xDB, 0xE1)
DARK_GRAY  = RGBColor(0x5D, 0x6D, 0x7E)
GREEN      = RGBColor(0x27, 0xAE, 0x60)
RED        = RGBColor(0xE7, 0x4C, 0x3C)
TEAL       = RGBColor(0x1A, 0xBC, 0x9C)
PURPLE     = RGBColor(0x8E, 0x44, 0xAD)
CARD_BG    = RGBColor(0x1A, 0x2F, 0x55)

def add_bg(slide, color=NAVY):
    bg = slide.background; fill = bg.fill; fill.solid(); fill.fore_color.rgb = color

def add_shape(slide, left, top, w, h, fill_color, border_color=None, border_w=Pt(0), shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, left, top, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = fill_color
    if border_color: shape.line.color.rgb = border_color; shape.line.width = border_w
    else: shape.line.fill.background()
    shape.shadow.inherit = False; return shape

def add_rect(slide, left, top, w, h, fill_color, border_color=None, border_w=Pt(0)):
    return add_shape(slide, left, top, w, h, fill_color, border_color, border_w, MSO_SHAPE.RECTANGLE)

def add_text_box(slide, left, top, w, h, text, font_size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, w, h)
    tf = txBox.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.font.size = Pt(font_size)
    p.font.color.rgb = color; p.font.bold = bold; p.font.name = "Microsoft JhengHei"; p.alignment = alignment
    return txBox

def add_multiline_text(slide, left, top, w, h, lines, line_spacing=1.3):
    txBox = slide.shapes.add_textbox(left, top, w, h)
    tf = txBox.text_frame; tf.word_wrap = True
    for i, (txt, c, b, fs) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt; p.font.size = Pt(fs); p.font.color.rgb = c; p.font.bold = b
        p.font.name = "Microsoft JhengHei"; p.space_after = Pt(fs * 0.3); p.line_spacing = Pt(fs * line_spacing)
    return txBox

def set_shape_text(shape, text, font_size=11, color=WHITE, bold=False, alignment=PP_ALIGN.CENTER):
    tf = shape.text_frame; tf.word_wrap = True; tf.paragraphs[0].alignment = alignment
    p = tf.paragraphs[0]; p.text = text; p.font.size = Pt(font_size); p.font.color.rgb = color; p.font.bold = bold; p.font.name = "Microsoft JhengHei"

# 載入原始 V2 檔案
input_path = r"d:\JavaDO\總結報告_v2.pptx"
prs = Presentation(input_path)
SLIDE_W = prs.slide_width
TOTAL_SLIDES = len(prs.slides) + 5  # 原 17 頁 + 5 頁補充

def add_slide_number(slide, num, total):
    add_text_box(slide, Inches(12.3), Inches(7.05), Inches(1), Inches(0.4), f"補充 - {num}", 9, DARK_GRAY, False, PP_ALIGN.RIGHT)

# ============================================================
# SUPPLEMENTARY SLIDE 1: 補充篇標題 (Slide 18)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(1.5), Inches(3.0), Inches(10), Inches(1.2), "補充篇：進階基礎設施與防護機制", 36, WHITE, True, PP_ALIGN.CENTER)
add_text_box(slide, Inches(1.5), Inches(4.0), Inches(10), Inches(0.8), "實體架構 / 護欄時序 / SBOM供應鏈 / ZTA與自動化紅隊", 18, LIGHT_BLUE, False, PP_ALIGN.CENTER)
add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)

# ============================================================
# SUPPLEMENTARY SLIDE 2: 硬體實體架構與零信任網路 (Slide 19)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "補充 01：硬體實體架構與零信任網路 (對應資安架構圖)", 26, WHITE, True)
add_text_box(slide, Inches(0.8), Inches(0.85), Inches(11), Inches(0.4), "導入 Kubernetes (K8s) 叢集、PAM 特權帳號管理與 PQ Tunnel (NIST SP 800-207 ZTA)", 12, LIGHT_GRAY, False)

# Network Layer
add_shape(slide, Inches(0.5), Inches(1.4), Inches(12.3), Inches(1.0), DARK_BLUE, ACCENT_BLUE, Pt(1))
add_text_box(slide, Inches(0.7), Inches(1.5), Inches(3), Inches(0.3), "🌐 核心交換與零信任通道", 13, GOLD, True)
add_text_box(slide, Inches(0.7), Inches(1.8), Inches(11.5), Inches(0.5), 
             "• Core Switch (100Gbps / 400Gbps) 高速骨幹網路 (Ethernet)\n"
             "• 導入 PQ Tunnel (後量子密碼學) 零信任網路管理平台，落實 NIST SP 800-207 (ZTA 網路支柱)", 11, LIGHT_GRAY, False)

# Zones
zones = [
    ("資料清洗區", RED, "• 單向閘道器 (Diode) #1-2\n• 資料清洗節點 #1\n• 負責外網資料過濾、惡意程式掃描與格式正規化"),
    ("算力資源區 (K8s / TEE)", MID_BLUE, "• 導入 Kubernetes (K8s) 架構進行運算資源排程\n• Liquid-Cooled AMD MI355X / NVIDIA HGX B200 節點\n• 啟用 TEE 機密運算保護運行中模型權重"),
    ("存取控制區 (PAM)", ORANGE, "• 特權帳號管理節點 (PAM) #1-2\n• 網路管理節點 #1-2\n• 嚴格收斂並管控 K8s 與 Mgt 節點之最高維運權限"),
    ("資安監控區", TEAL, "• 日誌管理節點 #1-4 (Mgt Node)\n• 威脅情資管理節點 #1-2 (Security Mgt Node)\n• 介接 SOC 中心與自動化紅隊演練系統"),
]
for i, (title, color, desc) in enumerate(zones):
    x = Inches(0.5) + Inches(i * 3.15)
    add_shape(slide, x, Inches(2.7), Inches(3.0), Inches(4.2), CARD_BG, color, Pt(1.5))
    add_rect(slide, x, Inches(2.7), Inches(3.0), Inches(0.5), color)
    add_text_box(slide, x, Inches(2.8), Inches(3.0), Inches(0.3), title, 13, WHITE, True, PP_ALIGN.CENTER)
    add_multiline_text(slide, x + Inches(0.15), Inches(3.4), Inches(2.7), Inches(3.0), [(line, LIGHT_GRAY, False, 11) for line in desc.split("\n")], 1.5)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 1, 5)

# ============================================================
# SUPPLEMENTARY SLIDE 3: 護欄運作循序流程 (Slide 20)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "補充 02：護欄運作雙向循序流程 (對應護欄循序圖)", 26, WHITE, True)
add_text_box(slide, Inches(0.8), Inches(0.85), Inches(11), Inches(0.4), "護欄服務作為獨立攔截點，強制執行 Input/Output 雙向檢核與稽核留存", 12, LIGHT_GRAY, False)

add_shape(slide, Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.5), DARK_BLUE, ACCENT_BLUE, Pt(1.5))

box_actor = add_shape(slide, Inches(0.8), Inches(2.5), Inches(1.5), Inches(0.8), CARD_BG, WHITE, Pt(1))
set_shape_text(box_actor, "👤\nActor (User)", 12, WHITE, True)

box_agent = add_shape(slide, Inches(4.5), Inches(2.2), Inches(2.5), Inches(1.2), MID_BLUE, TEAL, Pt(2))
set_shape_text(box_agent, "代理層 (Agent Layer)", 14, WHITE, True)

box_llm = add_shape(slide, Inches(9.5), Inches(2.2), Inches(2.5), Inches(1.2), MID_BLUE, ORANGE, Pt(2))
set_shape_text(box_llm, "LLM 模型", 14, WHITE, True)

box_guard = add_shape(slide, Inches(4.0), Inches(4.5), Inches(3.5), Inches(1.5), RED, DARK_GRAY, Pt(0))
set_shape_text(box_guard, "護欄服務 (Guardrail)", 14, WHITE, True)

box_admin = add_shape(slide, Inches(10.5), Inches(4.8), Inches(1.5), Inches(0.8), CARD_BG, WHITE, Pt(1))
set_shape_text(box_admin, "🛠️\nAdmin", 12, WHITE, True)

box_log = add_shape(slide, Inches(4.75), Inches(6.3), Inches(2.0), Inches(0.6), PURPLE)
set_shape_text(box_log, "🛢️ Audit Log", 12, WHITE, True)

def draw_arrow_label(slide, x, y, w, h, text, color=GOLD, shape_type=MSO_SHAPE.RIGHT_ARROW):
    arr = slide.shapes.add_shape(shape_type, x, y, w, h)
    arr.fill.solid(); arr.fill.fore_color.rgb = color; arr.line.fill.background()
    add_text_box(slide, x, y - Inches(0.25), w + Inches(1), Inches(0.3), text, 10, color, True, PP_ALIGN.CENTER)

draw_arrow_label(slide, Inches(2.4), Inches(2.8), Inches(1.9), Inches(0.15), "1. Request (用戶請求)", GOLD, MSO_SHAPE.RIGHT_ARROW)

draw_arrow_label(slide, Inches(4.6), Inches(3.5), Inches(0.15), Inches(0.9), "1.1 輸入", TEAL, MSO_SHAPE.DOWN_ARROW)
draw_arrow_label(slide, Inches(5.2), Inches(3.5), Inches(0.15), Inches(0.9), "1.2 回應(輸入確認)", TEAL, MSO_SHAPE.UP_ARROW)

draw_arrow_label(slide, Inches(7.2), Inches(2.5), Inches(2.1), Inches(0.15), "1.3 輸入 (安全 Prompt)", ORANGE, MSO_SHAPE.RIGHT_ARROW)
draw_arrow_label(slide, Inches(7.2), Inches(3.0), Inches(2.1), Inches(0.15), "2.1 回傳回應", ORANGE, MSO_SHAPE.LEFT_ARROW)

draw_arrow_label(slide, Inches(6.3), Inches(3.5), Inches(0.15), Inches(0.9), "2.2 輸出", TEAL, MSO_SHAPE.DOWN_ARROW)
draw_arrow_label(slide, Inches(6.9), Inches(3.5), Inches(0.15), Inches(0.9), "2.3 回應(輸出確認)", TEAL, MSO_SHAPE.UP_ARROW)

draw_arrow_label(slide, Inches(7.7), Inches(5.1), Inches(2.6), Inches(0.15), "0. 管理 System Prompt (政策碼)", ACCENT_BLUE, MSO_SHAPE.LEFT_ARROW)

draw_arrow_label(slide, Inches(5.65), Inches(6.0), Inches(0.2), Inches(0.25), "單向稽核", LIGHT_GRAY, MSO_SHAPE.DOWN_ARROW)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 2, 5)

# ============================================================
# SUPPLEMENTARY SLIDE 4: SBOM 供應鏈管控與 MCP (Slide 21)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "補充 03：SBOM 供應鏈管控與 MCP 零信任授權", 26, WHITE, True)
add_text_box(slide, Inches(0.8), Inches(0.85), Inches(11), Inches(0.4), "對齊 NIST SP 800-218A SSDF 規範，從模型到工具建立完整信任鏈", 12, LIGHT_GRAY, False)

add_shape(slide, Inches(0.5), Inches(1.5), Inches(5.8), Inches(5.4), CARD_BG, TEAL, Pt(2))
add_text_box(slide, Inches(0.7), Inches(1.7), Inches(5), Inches(0.3), "📦 SBOM (軟體料件清單) 管控", 15, GOLD, True)
sbom_desc = [
    ("基礎模型層 (L1)", ACCENT_BLUE, True, 13),
    ("• 模型權重下載後強制要求提供 SBOM。\n• 進行雜湊值 (Hash) 比對與數位簽章驗證。\n• 確認無遭竄改或植入後門方可掛載至算力區。", LIGHT_GRAY, False, 11),
    ("MCP 工具層 (L4)", TEAL, True, 13),
    ("• MCP 伺服器與外掛套件上架前，必須提交原始碼與相依套件 SBOM。\n• 經過自動化弱點掃描 (SCA) 排除已知 CVE 漏洞。\n• 於 Portal 上架時由資安官進行雙重確認。", LIGHT_GRAY, False, 11),
    ("稽核留存", PURPLE, True, 13),
    ("• SBOM 紀錄一併綁定至 8 欄位稽核日誌。\n• 落實全鏈路 Traceability (可追溯性)。", LIGHT_GRAY, False, 11)
]
add_multiline_text(slide, Inches(0.7), Inches(2.2), Inches(5.4), Inches(4.5), sbom_desc, 1.5)

add_shape(slide, Inches(6.8), Inches(1.5), Inches(6.0), Inches(5.4), CARD_BG, ORANGE, Pt(2))
add_text_box(slide, Inches(7.0), Inches(1.7), Inches(5), Inches(0.3), "🔑 MCP 工具零信任授權 (mTLS)", 15, GOLD, True)
mcp_desc = [
    ("拒絕靜態金鑰", RED, True, 13),
    ("• 傳統架構常將 API Key 寫死於 Agent 或環境變數中，一旦外洩即造成橫向移動威脅。\n• 本專案全面禁止 MCP 伺服器使用靜態 API Key 認證。", LIGHT_GRAY, False, 11),
    ("短時效憑證 (Ephemeral Credentials)", ORANGE, True, 13),
    ("• Agent 呼叫外部武器介面、雷達資料庫或機敏系統時，必須向 Gateway 申請短時效 Token (例如 5 分鐘)。\n• 逾期自動失效，縮小攻擊窗口。", LIGHT_GRAY, False, 11),
    ("mTLS 雙向認證", TEAL, True, 13),
    ("• 結合 SPIFFE/SPIRE 簽發工作負載身分。\n• MCP 伺服器與 Agent 之間強制建立 mTLS 雙向加密通道。", LIGHT_GRAY, False, 11)
]
add_multiline_text(slide, Inches(7.0), Inches(2.2), Inches(5.6), Inches(4.5), mcp_desc, 1.5)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 3, 5)

# ============================================================
# SUPPLEMENTARY SLIDE 5: ZTA 映射與自動化紅隊 (Slide 22)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "補充 04：ZTA 七大支柱映射與 AI 自動化紅隊演練", 26, WHITE, True)
add_text_box(slide, Inches(0.8), Inches(0.85), Inches(11), Inches(0.4), "嚴格對齊國防部零信任架構，並導入動態 AIEC (AI Evaluation Criteria) 防禦指標", 12, LIGHT_GRAY, False)

add_shape(slide, Inches(0.5), Inches(1.4), Inches(12.3), Inches(2.8), CARD_BG, ACCENT_BLUE, Pt(2))
add_text_box(slide, Inches(0.7), Inches(1.5), Inches(6), Inches(0.3), "🛡️ 國防部零信任 ZTA 七大支柱對齊", 14, GOLD, True)
zta_pillars = [
    ("User (使用者)", "SSO/MFA + ABAC 動態屬性授權"), ("Device (裝置)", "端點 MDM / 內網設備健康度檢查"),
    ("Network (網路)", "PQ Tunnel 後量子通道 + 單閘 Diode"), ("App (應用程式)", "Harness 治理外框 + K8s 沙箱隔離"),
    ("Data (資料)", "DISA IL 分級 + RAG 隔離 + DLP 遮罩"), ("Visibility (可視化)", "8 欄位留痕 + SBOM + SOC 監控"),
    ("Automation (自動化)", "異常自動熔斷斷路器 (Circuit Breaker)")
]
for i, (p, d) in enumerate(zta_pillars):
    x = Inches(0.7) + Inches((i % 4) * 2.95)
    y = Inches(2.0) + Inches((i // 4) * 0.9)
    add_text_box(slide, x, y, Inches(2.8), Inches(0.3), "柱 " + str(i+1) + ": " + p, 12, ACCENT_BLUE, True)
    add_text_box(slide, x, y + Inches(0.25), Inches(2.8), Inches(0.4), d, 10, LIGHT_GRAY, False)

add_shape(slide, Inches(0.5), Inches(4.4), Inches(12.3), Inches(2.6), CARD_BG, RED, Pt(2))
add_text_box(slide, Inches(0.7), Inches(4.5), Inches(8), Inches(0.3), "🔴 AI 驅動之自動化紅隊演練 (Automated AI Red Teaming)", 14, GOLD, True)
add_text_box(slide, Inches(0.7), Inches(4.9), Inches(11.5), Inches(2.0),
             "改變傳統靜態演練，引入「對抗性 LLM」自動對 Agent 進行 Prompt Injection、記憶投毒與沙箱越權測試，動態衡量系統韌性。\n\n"
             "【AIEC 量化防禦指標】\n"
             "▸ 提示抗注入率 (Prompt Robustness)：目標 > 99.9% 攔截率\n"
             "▸ 防降密洩漏率 (Data Exfiltration Prevention)：100% 機敏資料阻斷\n"
             "▸ 沙箱與 MCP 逃逸防禦 (Sandbox Escape Resilience)：零容忍\n"
             "▸ 護欄攔截準確率 (Guardrail Precision)：每月抽樣重跑比對基準", 12, LIGHT_GRAY, False)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 4, 5)

# ─── Save ───
output_path = r"d:\JavaDO\總結報告_v4_補充版.pptx"
try:
    prs.save(output_path)
    print(f"✅ 總結報告 V4_補充版 已成功生成：{output_path}")
    print(f"   總頁數：{TOTAL_SLIDES} (保留原 17 頁，新增補充 5 頁)")
except Exception as e:
    print(f"寫入失敗，請確認檔案未被開啟: {e}")
