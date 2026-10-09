#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
總結報告_最終版.pptx — Opus 4.6 Review 微調腳本
策略：讀取 總結報告_v4_補充版.pptx (22 頁)
  - Slide 1~17: 完全不動
  - Slide 18: 補充過場頁 → 追加 4 張摘要卡片
  - Slide 19: 完全不動
  - Slide 20: 護欄循序 → 追加底部阻斷/通過處置文字
  - Slide 21~22: 完全不動
  - 新增 Slide 23: 全案總結論
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import sys

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# ─── Color Palette (與 V2 相同) ───
NAVY       = RGBColor(0x0B, 0x1D, 0x3A)
DARK_BLUE  = RGBColor(0x14, 0x2D, 0x5E)
MID_BLUE   = RGBColor(0x1E, 0x3A, 0x6E)
ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xC1)
LIGHT_BLUE = RGBColor(0x5D, 0xAE, 0xF5)
GOLD       = RGBColor(0xF4, 0xC7, 0x30)
ORANGE     = RGBColor(0xE8, 0x7C, 0x2A)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xD5, 0xDB, 0xE1)
DARK_GRAY  = RGBColor(0x5D, 0x6D, 0x7E)
GREEN      = RGBColor(0x27, 0xAE, 0x60)
RED        = RGBColor(0xE7, 0x4C, 0x3C)
TEAL       = RGBColor(0x1A, 0xBC, 0x9C)
PURPLE     = RGBColor(0x8E, 0x44, 0xAD)
CARD_BG    = RGBColor(0x1A, 0x2F, 0x55)

def add_bg(slide, color=NAVY):
    bg = slide.background; fill = bg.fill; fill.solid(); fill.fore_color.rgb = color

def add_shape(slide, left, top, w, h, fill_color, border_color=None, border_w=Pt(0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color; shape.line.width = border_w
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False; return shape

def add_rect(slide, left, top, w, h, fill_color, border_color=None, border_w=Pt(0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color; shape.line.width = border_w
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False; return shape

def add_text_box(slide, left, top, w, h, text, font_size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, w, h)
    tf = txBox.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.font.size = Pt(font_size)
    p.font.color.rgb = color; p.font.bold = bold; p.font.name = "Microsoft JhengHei"; p.alignment = alignment
    return txBox

def set_shape_text(shape, text, font_size=11, color=WHITE, bold=False, alignment=PP_ALIGN.CENTER):
    tf = shape.text_frame; tf.word_wrap = True; tf.paragraphs[0].alignment = alignment
    p = tf.paragraphs[0]; p.text = text; p.font.size = Pt(font_size)
    p.font.color.rgb = color; p.font.bold = bold; p.font.name = "Microsoft JhengHei"

def add_multiline(slide, left, top, w, h, lines, line_spacing=1.3):
    txBox = slide.shapes.add_textbox(left, top, w, h)
    tf = txBox.text_frame; tf.word_wrap = True
    for i, (txt, c, b, fs) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt; p.font.size = Pt(fs); p.font.color.rgb = c; p.font.bold = b
        p.font.name = "Microsoft JhengHei"; p.space_after = Pt(fs * 0.3)
        p.line_spacing = Pt(fs * line_spacing)
    return txBox


# ─── 載入 V4 ───
input_path = r"d:\JavaDO\總結報告_v4_補充版.pptx"
prs = Presentation(input_path)
SLIDE_W = prs.slide_width

print(f"已載入 {input_path}，共 {len(prs.slides)} 頁")

# ============================================================
# 微調 Slide 18 (index 17): 補充篇過場頁 — 追加 4 張摘要卡片
# ============================================================
slide18 = prs.slides[17]

cards = [
    ("🏗️ 硬體實體架構", ACCENT_BLUE,
     "K8s 容器排程 · PAM 特權管控\nPQ Tunnel 後量子零信任\nTEE 機密運算保護模型權重"),
    ("🔄 護欄雙向時序", TEAL,
     "Input/Output 獨立攔截點\n阻斷→替換警告 / 通過→完整回傳\nAdmin System Prompt 政策控制"),
    ("📦 SBOM + MCP 授權", ORANGE,
     "模型/工具 SBOM 料件清單驗證\nSCA 弱掃排除已知 CVE\nmTLS 短時效憑證 零信任授權"),
    ("🛡️ ZTA + AI 紅隊", RED,
     "國防部 ZTA 七大支柱全對齊\nAI 對抗性 LLM 自動化紅隊\nAIEC 量化防禦指標體系"),
]
for i, (title, color, desc) in enumerate(cards):
    x = Inches(0.5) + Inches(i * 3.2)
    card = add_shape(slide18, x, Inches(5.0), Inches(2.9), Inches(2.1), CARD_BG, color, Pt(1.5))
    badge = add_shape(slide18, x + Inches(0.05), Inches(5.1), Inches(2.8), Inches(0.45), color)
    set_shape_text(badge, title, 12, WHITE, True)
    add_multiline(slide18, x + Inches(0.1), Inches(5.65), Inches(2.7), Inches(1.3),
                  [(line, LIGHT_GRAY, False, 10) for line in desc.split("\n")], 1.5)

print("✓ Slide 18 微調完成 (追加 4 張摘要卡片)")


# ============================================================
# 微調 Slide 20 (index 19): 護欄循序 — 追加處置說明
# ============================================================
slide20 = prs.slides[19]

# 在底部追加阻斷/通過處置說明文字框 (不碰原有元素)
# 先加背景卡片
notice_bg = add_rect(slide20, Inches(0.5), Inches(7.0), Inches(12.3), Inches(0.4), CARD_BG, GOLD, Pt(1))
add_text_box(slide20, Inches(0.7), Inches(7.02), Inches(12), Inches(0.35),
             "⚠️ 阻斷 → 代理層將回應替換為系統警告字串後回傳使用者 │ ✅ 通過 → 完整回應回傳給使用者，完成推論請求",
             10, GOLD, True, PP_ALIGN.CENTER)

print("✓ Slide 20 微調完成 (追加阻斷/通過處置說明)")


# ============================================================
# 新增 Slide 23: 全案總結論
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6),
             "全案總結論 — 五大原則 × 四大補強 完整覆蓋", 28, WHITE, True)

# ── 五大原則 × 四大補強 覆蓋矩陣 ──
add_shape(slide, Inches(0.3), Inches(1.1), Inches(12.7), Inches(3.0), CARD_BG, GOLD, Pt(2))
add_text_box(slide, Inches(0.5), Inches(1.2), Inches(8), Inches(0.3),
             "五大原則 × 四大補強 覆蓋矩陣", 14, GOLD, True)

# Headers
principles = ["模型分層", "資料分級", "單一閘道", "層層防護", "全程留痕"]
badge_colors = [ACCENT_BLUE, TEAL, ORANGE, PURPLE, GREEN]
for i, (p, c) in enumerate(zip(principles, badge_colors)):
    x = Inches(2.8) + Inches(i * 2.0)
    hdr = add_shape(slide, x, Inches(1.55), Inches(1.85), Inches(0.38), c)
    set_shape_text(hdr, p, 10, WHITE, True)

# Row headers + cells
supplements = [
    ("K8s / PAM / PQ Tunnel", ACCENT_BLUE),
    ("護欄雙向時序", TEAL),
    ("SBOM / MCP mTLS", ORANGE),
    ("ZTA 七柱 / AI 紅隊", RED),
]
# Coverage matrix: which principle is covered by which supplement
# Each row = [模型分層, 資料分級, 單一閘道, 層層防護, 全程留痕]
matrix = [
    ["K8s GPU\n隔離排程", "IL 分級\n分區部署", "PQ Tunnel\n單一通道", "TEE 機密\n運算防護", "PAM 操作\n全程監錄"],
    ["L1~L5\n分層攔截", "密級比對\n先過濾後召回", "Gateway\n單一檢核", "Input/Output\n雙向護欄", "Audit Log\n單向寫入"],
    ["模型 SBOM\n簽章驗證", "工具依密級\n白名單控管", "mTLS\n集中授權", "SCA 弱掃\nCVE 排除", "料件清單\n全鏈追溯"],
    ["七柱對齊\nL1~L5 映射", "Data 支柱\nDISA IL", "Network 支柱\nDiode 閘道", "App 支柱\nHarness 沙箱", "Visibility\nSOC+AIEC"],
]
for row_i, ((row_title, row_color), cells) in enumerate(zip(supplements, matrix)):
    y = Inches(2.0) + Inches(row_i * 0.5)
    # Row header
    rh = add_rect(slide, Inches(0.5), y, Inches(2.2), Inches(0.45), row_color)
    set_shape_text(rh, row_title, 9, WHITE, True)
    # Cells
    for col_i, cell_text in enumerate(cells):
        x = Inches(2.8) + Inches(col_i * 2.0)
        bg = CARD_BG if row_i % 2 == 0 else DARK_BLUE
        cell = add_rect(slide, x, y, Inches(1.85), Inches(0.45), bg, badge_colors[col_i], Pt(0.5))
        set_shape_text(cell, cell_text.split("\n")[0], 8, LIGHT_GRAY, False)

# ── 決策建議 ──
add_shape(slide, Inches(0.3), Inches(4.3), Inches(6.0), Inches(2.8), CARD_BG, GREEN, Pt(2))
add_text_box(slide, Inches(0.5), Inches(4.4), Inches(5.5), Inches(0.3),
             "🎯 對長官的三項決策建議", 15, GOLD, True)
add_multiline(slide, Inches(0.5), Inches(4.9), Inches(5.5), Inches(2.0), [
    ("建議一：立即啟動 Gateway + PQ Tunnel 建置", GREEN, True, 12),
    ("以 Q1 完成最小可行版為目標，確保全院 AI 流量「單一閘道 + 後量子加密」雙保險上線。", LIGHT_GRAY, False, 10),
    ("", WHITE, False, 6),
    ("建議二：先導單位 PoC 雙軌並行", TEAL, True, 12),
    ("行政庶務 Agent 以開箱即用 Harness 快速落地；同步啟動武器系統 Agent 的 DeepAgents 客製化框架評估。", LIGHT_GRAY, False, 10),
    ("", WHITE, False, 6),
    ("建議三：建立 SBOM + AI 紅隊常態機制", ORANGE, True, 12),
    ("將 SBOM 料件清單驗證納入 CI/CD 流程；季度執行 AI 自動化紅隊演練，以 AIEC 指標量化防禦韌性。", LIGHT_GRAY, False, 10),
], 1.3)

# ── 全案交付成果 ──
add_shape(slide, Inches(6.5), Inches(4.3), Inches(6.5), Inches(2.8), CARD_BG, ACCENT_BLUE, Pt(2))
add_text_box(slide, Inches(6.7), Inches(4.4), Inches(6), Inches(0.3),
             "📋 全案 23 頁核心交付成果", 15, GOLD, True)
deliverables = [
    ("主體篇 17 頁", ACCENT_BLUE, True, 12),
    ("五大原則 · 六節點縱深 · Harness 雙軌\n11步資料流 · Guardrail 五類佈設\n國際標準 29 項 · SOC 熔斷\n端到端 16 項檢核 · Q1~Q4 Roadmap\n風險矩陣 · KPI · 行動建議", LIGHT_GRAY, False, 10),
    ("", WHITE, False, 4),
    ("補充篇 6 頁", ORANGE, True, 12),
    ("K8s/PAM/PQ Tunnel 實體架構\n護欄雙向 Input/Output 時序\nSBOM 供應鏈 + MCP mTLS 零信任\nZTA 七大支柱 + AI 自動化紅隊\n五大原則×四大補強 覆蓋矩陣", LIGHT_GRAY, False, 10),
]
add_multiline(slide, Inches(6.7), Inches(4.9), Inches(6.0), Inches(2.0), deliverables, 1.25)

# ── 底部：五大原則徽章 (與封面呼應) ──
for i, (p, c) in enumerate(zip(principles, badge_colors)):
    s = add_shape(slide, Inches(0.8) + Inches(i * 2.5), Inches(7.35), Inches(2.2), Inches(0.0), c)
# 改為底部金色線
add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)

# 底部原則徽章放在金線上方
for i, (p, c) in enumerate(zip(principles, badge_colors)):
    s = add_shape(slide, Inches(1.6) + Inches(i * 2.1), Inches(7.15), Inches(1.8), Inches(0.25), c)
    set_shape_text(s, p, 10, WHITE, True)

add_text_box(slide, Inches(12.0), Inches(7.05), Inches(1.3), Inches(0.4),
             f"23 / 23", 9, DARK_GRAY, False, PP_ALIGN.RIGHT)

print("✓ Slide 23 新增完成 (全案總結論)")


# ─── Save ───
output_path = r"d:\JavaDO\總結報告_最終版.pptx"
try:
    prs.save(output_path)
    print(f"\n✅ 總結報告_最終版已成功生成：{output_path}")
    print(f"   共 {len(prs.slides)} 頁投影片")
    print("   微調項目：")
    print("   ✓ Slide 18：追加 4 張補強項目摘要卡片")
    print("   ✓ Slide 20：追加護欄阻斷/通過處置說明")
    print("   ✓ Slide 23：新增全案總結論 (五大原則×四大補強覆蓋矩陣 + 決策建議)")
except Exception as e:
    print(f"❌ 寫入失敗: {e}")
