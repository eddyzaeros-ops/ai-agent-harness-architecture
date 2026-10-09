#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
總結報告.pptx — AI 服務鏈總結報告生成腳本 v2
五大原則：模型分層、資料分級、單一閘道、層層防護、全程留痕
修正：
1. 開箱即用 Harness = Claude Code / OpenAI Codex / Google Antigravity 2.0
   客製化 Harness = LangChain DeepAgents
2. 國際標準全面補充：NIST RMF 1.0 / AI 600-1 / SP 800-53 R5 / OWASP AISVS / OWASP Top 10 / MITRE ATLAS 等
3. 對齊整合提報版 20260906 內容
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os, sys

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# ─── Color Palette ───
NAVY       = RGBColor(0x0B, 0x1D, 0x3A)
DARK_BLUE  = RGBColor(0x14, 0x2D, 0x5E)
MID_BLUE   = RGBColor(0x1E, 0x3A, 0x6E)
ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xC1)
LIGHT_BLUE = RGBColor(0x5D, 0xAE, 0xF5)
GOLD       = RGBColor(0xF4, 0xC7, 0x30)
ORANGE     = RGBColor(0xE8, 0x7C, 0x2A)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xD5, 0xDB, 0xE1)
DARK_GRAY  = RGBColor(0x5D, 0x6D, 0x7E)
GREEN      = RGBColor(0x27, 0xAE, 0x60)
RED        = RGBColor(0xE7, 0x4C, 0x3C)
TEAL       = RGBColor(0x1A, 0xBC, 0x9C)
PURPLE     = RGBColor(0x8E, 0x44, 0xAD)
SOFT_BG    = RGBColor(0xF0, 0xF3, 0xF7)
CARD_BG    = RGBColor(0x1A, 0x2F, 0x55)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

def add_bg(slide, color=NAVY):
    bg = slide.background; fill = bg.fill; fill.solid(); fill.fore_color.rgb = color

def add_shape(slide, left, top, w, h, fill_color, border_color=None, border_w=Pt(0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = fill_color
    if border_color: shape.line.color.rgb = border_color; shape.line.width = border_w
    else: shape.line.fill.background()
    shape.shadow.inherit = False; return shape

def add_rect(slide, left, top, w, h, fill_color, border_color=None, border_w=Pt(0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = fill_color
    if border_color: shape.line.color.rgb = border_color; shape.line.width = border_w
    else: shape.line.fill.background()
    shape.shadow.inherit = False; return shape

def add_text_box(slide, left, top, w, h, text, font_size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Microsoft JhengHei"):
    txBox = slide.shapes.add_textbox(left, top, w, h)
    tf = txBox.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.font.size = Pt(font_size)
    p.font.color.rgb = color; p.font.bold = bold; p.font.name = font_name; p.alignment = alignment
    return txBox

def add_multiline_text(slide, left, top, w, h, lines, font_size=12, color=WHITE, bold=False, line_spacing=1.3, font_name="Microsoft JhengHei"):
    txBox = slide.shapes.add_textbox(left, top, w, h)
    tf = txBox.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        if isinstance(line, str): txt, c, b, fs = line, color, bold, font_size
        elif len(line) == 2: txt, c = line; b = bold; fs = font_size
        elif len(line) == 3: txt, c, b = line; fs = font_size
        else: txt, c, b, fs = line
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt; p.font.size = Pt(fs); p.font.color.rgb = c; p.font.bold = b
        p.font.name = font_name; p.space_after = Pt(font_size * 0.3); p.line_spacing = Pt(font_size * line_spacing)
    return txBox

def add_arrow_shape(slide, left, top, w, h, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = fill_color; shape.line.fill.background(); shape.shadow.inherit = False; return shape

def add_down_arrow(slide, left, top, w, h, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = fill_color; shape.line.fill.background(); shape.shadow.inherit = False; return shape

def set_shape_text(shape, text, font_size=11, color=WHITE, bold=False, alignment=PP_ALIGN.CENTER, font_name="Microsoft JhengHei"):
    tf = shape.text_frame; tf.word_wrap = True; tf.paragraphs[0].alignment = alignment
    p = tf.paragraphs[0]; p.text = text; p.font.size = Pt(font_size); p.font.color.rgb = color; p.font.bold = bold; p.font.name = font_name; return tf

def add_slide_number(slide, num, total):
    add_text_box(slide, Inches(12.3), Inches(7.05), Inches(1), Inches(0.4), f"{num} / {total}", 9, DARK_GRAY, False, PP_ALIGN.RIGHT)

TOTAL_SLIDES = 17

# ============================================================
# SLIDE 1: 封面
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(1.5), Inches(1.5), Inches(10), Inches(1.2), "國防 AI 服務鏈 總結報告", 42, WHITE, True, PP_ALIGN.CENTER)
add_text_box(slide, Inches(1.5), Inches(2.7), Inches(10), Inches(0.8), "AI Service Chain — 架構研析 · 安全治理 · 建置藍圖", 22, LIGHT_BLUE, False, PP_ALIGN.CENTER)

principles = ["模型分層", "資料分級", "單一閘道", "層層防護", "全程留痕"]
badge_colors = [ACCENT_BLUE, TEAL, ORANGE, PURPLE, GREEN]
for i, (p, c) in enumerate(zip(principles, badge_colors)):
    s = add_shape(slide, Inches(2.2) + Inches(i * 1.85), Inches(4.0), Inches(1.6), Inches(0.55), c)
    set_shape_text(s, p, 14, WHITE, True)

add_text_box(slide, Inches(1.5), Inches(5.3), Inches(10), Inches(0.5),
             "研析引擎：Claude Opus 4.6 (Thinking) │ 報告日期：2026-09-13", 14, DARK_GRAY, False, PP_ALIGN.CENTER)
add_text_box(slide, Inches(1.5), Inches(5.8), Inches(10), Inches(0.5),
             "涵蓋範圍：Agent 開發流程 · 使用者操作流程 · AI 資安防護 · Roadmap 效益評估", 13, LIGHT_GRAY, False, PP_ALIGN.CENTER)
add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 1, TOTAL_SLIDES)

# ============================================================
# SLIDE 2: 目錄
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7), "報告目錄 Table of Contents", 28, WHITE, True)

sections = [
    ("01", "五大原則總覽", "模型分層 · 資料分級 · 單一閘道 · 層層防護 · 全程留痕", ACCENT_BLUE),
    ("02", "AI 服務鏈全景架構", "六節點縱深防禦：防火牆→Portal→Gateway→算力→沙箱→Guardrail", ACCENT_BLUE),
    ("03", "Agent 開發者 — 行政庶務 Agent", "開箱即用 Harness：Claude Code / Codex / Antigravity 2.0", TEAL),
    ("04", "Agent 開發者 — 武器系統 Agent", "客製化 Harness：LangChain DeepAgents · 三重沙箱 · HITL", TEAL),
    ("05", "Agent 開發者 — 管控與軌跡", "Harness 治理外框 · CI/CD · 可觀測性 · 8欄位稽核紀錄", TEAL),
    ("06", "Agent 使用者 — Portal 操作流程", "Portal→Gateway 11步資料流 · RBAC+ABAC · HITL 審批", ORANGE),
    ("07", "Agent 使用者 — 管控與軌跡", "Portal/Gateway 雙管控 · 不可否認紀錄 · fail-closed", ORANGE),
    ("08", "資安防護：單閘 Diode · DLP", "單向閘道 · 資料外洩防護 · PII遮罩 · 落點路由", PURPLE),
    ("09", "資安防護：資料治理 · Guardrail", "五階資料管線 · 五類 Guardrail · 處置矩陣", PURPLE),
    ("10", "資安防護：AI 國際標準全覽", "ISO 42001 · NIST AI RMF · OWASP AISVS · MITRE ATLAS 等", PURPLE),
    ("11", "資安防護：SOC 監控與熔斷", "SIEM · 異常偵測 · 斷路器 · 事件回應 · 紅隊演練", PURPLE),
    ("12", "AI 服務鏈 Roadmap", "四階段建置：Q1基礎→Q2先導→Q3串聯→Q4全院合規", GREEN),
    ("13", "效益評估與 KPI", "100%閘道覆蓋 · 0件出境事件 · 影子AI收斂", GREEN),
    ("14", "結論與下一步行動", "關鍵成果 · 決策建議 · 後續規劃", GOLD),
]
y = Inches(1.3)
for i, (num, title, desc, color) in enumerate(sections):
    col_x = Inches(0.5) if i < 7 else Inches(6.8)
    row_y = y + Inches((i if i < 7 else i - 7) * 0.78)
    ns = add_shape(slide, col_x, row_y, Inches(0.55), Inches(0.55), color)
    set_shape_text(ns, num, 13, WHITE, True)
    add_text_box(slide, col_x + Inches(0.7), row_y, Inches(5), Inches(0.3), title, 13, WHITE, True)
    add_text_box(slide, col_x + Inches(0.7), row_y + Inches(0.28), Inches(5), Inches(0.3), desc, 9, LIGHT_GRAY, False)
add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 2, TOTAL_SLIDES)

# ============================================================
# SLIDE 3: 五大原則總覽
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(8), Inches(0.6), "01  五大原則總覽", 26, WHITE, True)
add_text_box(slide, Inches(0.8), Inches(0.85), Inches(11), Inches(0.4),
             "所有設計決策均錨定於五大安全治理原則，貫穿 AI 服務鏈全生命週期", 13, LIGHT_GRAY, False)

principle_data = [
    ("模型分層", ACCENT_BLUE, "Model Tiering",
     "• L1 通用基礎模型\n• L2 領域 DSLM\n• L3 Agent 代理\n• L4 Skill / 工具\n• L5 領域知識庫\n• T1~T3 地端 / T4 雲端路由"),
    ("資料分級", TEAL, "Data Classification",
     "• 公開 IL-2（TLP:CLEAR）\n• 營業秘密 IL-4（TLP:AMBER）\n• 核心營業秘密 IL-5\n• 機密 IL-6（TLP:RED）\n• C-I-A 三元組評定\n• 標籤繼承至 RAG 切片"),
    ("單一閘道", ORANGE, "Single Gateway",
     "• Portal＝單一入口\n• Gateway＝單一出口\n• 金鑰僅存 Gateway 密鑰庫\n• 防火牆封鎖 API 直連\n• OPA/Rego 三維裁決\n• fail-closed 預設拒絕"),
    ("層層防護", PURPLE, "Defense in Depth",
     "• Harness 治理外框\n• 六節點縱深防禦\n• 五類 Guardrail 佈設\n• L0~L3 四級沙箱隔離\n• HITL interrupt() 閘門\n• 對齊 ISO 42001 AIMS"),
    ("全程留痕", GREEN, "Full Auditability",
     "• 8 欄位不可否認紀錄\n• append-only 防竄改 ≥ 1 年\n• 全鏈路 trace 可觀測性\n• SOC 即時告警介接\n• 異常熔斷斷路器\n• 稽核 + 紅隊演練佐證"),
]
for i, (title, color, eng, details) in enumerate(principle_data):
    x = Inches(0.4) + Inches(i * 2.55)
    card = add_shape(slide, x, Inches(1.5), Inches(2.35), Inches(5.5), CARD_BG, color, Pt(2))
    badge = add_shape(slide, x + Inches(0.15), Inches(1.7), Inches(2.05), Inches(0.65), color)
    set_shape_text(badge, title, 18, WHITE, True)
    add_text_box(slide, x + Inches(0.15), Inches(2.45), Inches(2.05), Inches(0.35), eng, 10, LIGHT_BLUE, False, PP_ALIGN.CENTER)
    add_multiline_text(slide, x + Inches(0.2), Inches(2.9), Inches(2.0), Inches(3.8),
                       [(line, LIGHT_GRAY, False, 10) for line in details.split("\n")], 10, LIGHT_GRAY, False, 1.6)
add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 3, TOTAL_SLIDES)

# ============================================================
# SLIDE 4: AI 服務鏈全景架構（對齊整合提報版 Slide 5）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "02  AI 服務鏈全景架構 — 六節點縱深防禦", 26, WHITE, True)
add_text_box(slide, Inches(0.8), Inches(0.85), Inches(12), Inches(0.4),
             "防火牆 → Portal → Gateway → 核心算力 → 沙箱/Container → 輸出 Guardrail　每一節點皆為獨立管控點", 12, LIGHT_GRAY, False)

nodes = [
    ("01 邊界防火牆", "NGFW+IPS\n預設拒絕白名單放行\n南北向微分段\n出向封鎖 API 直連\n跨域不開雙向通道", "NIST SP 800-41\nA 級防護基準", ACCENT_BLUE),
    ("02 使用者 Portal", "SSO+MFA 雙因素\n會話綁定部門\n上傳掃描+分級標示\n提示前置 DLP 檢查\n禁跨部門歷史檢視", "ISO 27001 A.5.15\nOWASP AISVS C1", TEAL),
    ("03 AI Gateway", "OPA/Rego 三維裁決\n分級路由: 敏感→地端\n注入偵測+速率配額\n全量留跡: 提示/模型\n金鑰集中+fail-closed", "OWASP LLM01/10\nISO 42001 A.6", ORANGE),
    ("04 核心算力", "四層模型堆疊\n模型權重簽章+SBOM\n部門 GPU 配額隔離\n離線評測後上線\n不以敏感資料微調", "NIST AI RMF\nISO 42001", MID_BLUE),
    ("05 沙箱/容器", "每次呼叫獨立容器\n用後即銷毀\n預設無網路+唯讀\nAppArmor+cgroup\nMCP 自建自審白名單", "OWASP AISVS C10\nLLM06 過度代理", PURPLE),
    ("06 輸出 Guardrail", "敏感資訊偵測遮蔽\n越級阻斷+封鎖會話\nHITL interrupt() 核可\n輸出加註分級標籤\n可追溯浮水印", "OWASP LLM02\nISO 42001 A.9", RED),
]
for i, (title, desc, ref, color) in enumerate(nodes):
    x = Inches(0.2) + Inches(i * 2.15)
    card = add_shape(slide, x, Inches(1.4), Inches(2.0), Inches(4.5), CARD_BG, color, Pt(1.5))
    badge = add_shape(slide, x + Inches(0.08), Inches(1.52), Inches(1.84), Inches(0.5), color)
    set_shape_text(badge, title, 10, WHITE, True)
    add_multiline_text(slide, x + Inches(0.1), Inches(2.15), Inches(1.8), Inches(2.6),
                       [(line, LIGHT_GRAY, False, 9) for line in desc.split("\n")], 9, LIGHT_GRAY, False, 1.5)
    add_multiline_text(slide, x + Inches(0.1), Inches(4.9), Inches(1.8), Inches(0.8),
                       [(line, color, True, 8) for line in ref.split("\n")], 8, color, True, 1.4)
    if i < 5:
        add_arrow_shape(slide, x + Inches(2.03), Inches(3.3), Inches(0.2), Inches(0.25), GOLD)

# Harness bar
add_rect(slide, Inches(0.2), Inches(6.1), Inches(12.9), Inches(0.45), CARD_BG, GOLD, Pt(1.5))
add_text_box(slide, Inches(0.5), Inches(6.14), Inches(12.3), Inches(0.35),
             "Harness 治理外框（橫跨五層）── 模型負責「想」，Harness 負責「准不准做」── 工具白名單 · 重大動作先問人 · 用量上限 · 全程留紀錄",
             10, GOLD, True, PP_ALIGN.CENTER)

# Bottom note
add_text_box(slide, Inches(0.5), Inches(6.65), Inches(12), Inches(0.6),
             "回應路徑：核心算力/沙箱 → 輸出Guardrail → Gateway留跡 → Portal → 使用者。任一節點命中規則即阻斷、留跡並告警，不得由使用者自行重試繞過。",
             9, LIGHT_GRAY, False, PP_ALIGN.CENTER)
add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 4, TOTAL_SLIDES)

# ============================================================
# SLIDE 5: Agent 開發者流程 — 行政庶務 Agent（開箱即用 Harness）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "03  Agent 開發者 — 行政庶務 Agent（開箱即用 Harness）", 24, WHITE, True)
add_text_box(slide, Inches(0.8), Inches(0.85), Inches(11), Inches(0.4),
             "適用場景：公文撰擬、會議摘要、知識問答、報表產製 │ 安全等級：DISA IL-2 ~ IL-4", 12, LIGHT_GRAY, False)

# Three OOB Harness options
add_text_box(slide, Inches(0.5), Inches(1.3), Inches(12), Inches(0.4), "開箱即用 Harness 選項 — 自帶完整治理框架，快速部署", 14, GOLD, True)
harness_opts = [
    ("Claude Code", "Anthropic", ACCENT_BLUE, [
        "• 內建 Harness 治理外框",
        "• 自帶 HITL 審批機制",
        "• 檔案系統權限控制",
        "• 工具白名單管理",
        "• 沙箱隔離執行環境",
        "• 即裝即用・低學習曲線",
    ]),
    ("OpenAI Codex", "OpenAI", TEAL, [
        "• Codex Agent 執行環境",
        "• 內建程式碼沙箱隔離",
        "• 工具呼叫標準化協議",
        "• 自動化任務拆解引擎",
        "• 輸入/輸出護欄預置",
        "• 雲端/地端模型切換",
    ]),
    ("Antigravity 2.0", "Google DeepMind", ORANGE, [
        "• 子 Agent 委派 + 隔離",
        "• 虛擬檔案系統 (VFS)",
        "• 自帶可觀測性追蹤",
        "• 記憶與上下文管理",
        "• 多模態輸入支援",
        "• 企業級安全防護",
    ]),
]
for i, (name, vendor, color, items) in enumerate(harness_opts):
    x = Inches(0.3) + Inches(i * 4.3)
    card = add_shape(slide, x, Inches(1.8), Inches(4.1), Inches(3.0), CARD_BG, color, Pt(1.5))
    badge = add_shape(slide, x + Inches(0.1), Inches(1.95), Inches(3.9), Inches(0.5), color)
    set_shape_text(badge, f"{name}  ({vendor})", 13, WHITE, True)
    add_multiline_text(slide, x + Inches(0.15), Inches(2.6), Inches(3.8), Inches(2.0),
                       [(item, LIGHT_GRAY, False, 10) for item in items], 10, LIGHT_GRAY, False, 1.4)

# Flow steps
add_text_box(slide, Inches(0.5), Inches(5.0), Inches(12), Inches(0.35), "開發流程 — 選用開箱即用 Harness 快速部署", 13, WHITE, True)
flow_steps = [
    ("① 需求定義", "場景確認 · IL等級"),
    ("② Harness 選型", "Claude/Codex/AGY"),
    ("③ 工具配置", "MCP自建自審·白名單"),
    ("④ 權限 + HITL", "Deny-First · 審批點"),
    ("⑤ 安全測試", "越獄/注入/權限測試"),
    ("⑥ 容器化部署", "CI/CD · 映像掃描"),
]
for i, (title, desc) in enumerate(flow_steps):
    x = Inches(0.3) + Inches(i * 2.15)
    s = add_shape(slide, x, Inches(5.45), Inches(1.95), Inches(0.75), CARD_BG, TEAL, Pt(1))
    tf = s.text_frame; tf.word_wrap = True; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(10); p.font.color.rgb = TEAL; p.font.bold = True; p.font.name = "Microsoft JhengHei"
    p2 = tf.add_paragraph(); p2.text = desc; p2.font.size = Pt(9); p2.font.color.rgb = LIGHT_GRAY; p2.font.name = "Microsoft JhengHei"; p2.alignment = PP_ALIGN.CENTER
    if i < 5:
        add_arrow_shape(slide, x + Inches(2.0), Inches(5.7), Inches(0.22), Inches(0.25), TEAL)

# Shared controls
add_rect(slide, Inches(0.3), Inches(6.4), Inches(12.7), Inches(0.9), CARD_BG, ACCENT_BLUE, Pt(1))
add_text_box(slide, Inches(0.5), Inches(6.45), Inches(12.3), Inches(0.75),
             "共通管控要求 ▸ Harness 以基礎映像強制注入，比對失敗即拒絕啟動 ▸ MCP 伺服器一律自建自審，禁止直連外部未審伺服器 "
             "▸ 工具採白名單 ▸ 高衝擊動作強制 HITL ▸ 全量 append-only 稽核日誌",
             9, LIGHT_GRAY, False, PP_ALIGN.LEFT)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 5, TOTAL_SLIDES)

# ============================================================
# SLIDE 6: Agent 開發者流程 — 武器系統 Agent（客製化 Harness）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "04  Agent 開發者 — 武器系統 Agent（客製化 Harness）", 24, WHITE, True)
add_text_box(slide, Inches(0.8), Inches(0.85), Inches(11), Inches(0.4),
             "適用場景：戰術決策支援、情資分析、武器效能評估 │ 安全等級：DISA IL-5 ~ IL-6 │ LangChain DeepAgents 客製化", 12, LIGHT_GRAY, False)

# Left: Custom Harness Architecture
add_shape(slide, Inches(0.3), Inches(1.3), Inches(6.2), Inches(5.8), CARD_BG, RED, Pt(2))
add_text_box(slide, Inches(0.5), Inches(1.4), Inches(5.5), Inches(0.4), "LangChain DeepAgents 客製化 Harness", 15, GOLD, True)

weapon_steps = [
    ("① 三層堆疊客製", "Runtime(LangGraph) → Framework(LangChain) → Harness(DeepAgents)\n各層獨立配置、安全策略疊加、Middleware 鉤子注入"),
    ("② Context Engineering", "CompositeBackend (VFS) 大檔案卸載磁碟/記憶體\nread_file(offset,limit) 精準讀取，避免上下文膨脹"),
    ("③ 強化權限+沙箱", "L0 語言層→L1 程序層→L2 容器→L3 微虛擬機 四級隔離\nDeny-First + Catch-All Deny 絕對防線 · 即拋式執行環境"),
    ("④ 多層 HITL 審批", "武器決策→指揮官→法律官 鏈式審批\ninterrupt_on 多點掛載 · Checkpointer 斷點恢復"),
    ("⑤ 全規範合規驗證", "ISO 42001 · NIST AI RMF · OWASP AISVS · MITRE ATLAS\nAgent/Skills 安全基準全覆蓋 · 紅隊演練門檻"),
]
for i, (title, desc) in enumerate(weapon_steps):
    y = Inches(2.0) + Inches(i * 1.0)
    add_text_box(slide, Inches(0.5), y, Inches(2.0), Inches(0.3), title, 11, ORANGE, True)
    add_text_box(slide, Inches(2.5), y, Inches(3.8), Inches(0.85), desc, 9, LIGHT_GRAY, False)

# Right: Comparison table
add_shape(slide, Inches(6.8), Inches(1.3), Inches(6.2), Inches(5.8), CARD_BG, ACCENT_BLUE, Pt(2))
add_text_box(slide, Inches(7.0), Inches(1.4), Inches(5), Inches(0.4), "Harness 選型對照表", 15, GOLD, True)

headers = ["比較維度", "開箱即用 Harness", "客製化 Harness"]
for i, h in enumerate(headers):
    x = Inches(7.0) + Inches(i * 1.95)
    hdr = add_shape(slide, x, Inches(2.0), Inches(1.85), Inches(0.42), ACCENT_BLUE)
    set_shape_text(hdr, h, 10, WHITE, True)

rows = [
    ("代表工具", "Claude Code\nCodex · AGY 2.0", "LangChain\nDeepAgents v0.7"),
    ("適用場景", "行政庶務 Agent", "武器系統 Agent"),
    ("資料密級", "IL-2 ~ IL-4", "IL-5 ~ IL-6"),
    ("權限模式", "內建 Deny-First", "多層 ACL+四級沙箱"),
    ("HITL 審批", "單層審批", "多層鏈式審批"),
    ("模型選用", "雲端/地端通用", "僅限地端隔離模型"),
    ("執行隔離", "容器化沙箱", "gVisor/Kata/Firecracker"),
    ("MCP 治理", "標準工具白名單", "自建自審+能力憑證"),
    ("合規覆蓋", "基礎合規", "全規範+紅隊演練"),
    ("開發週期", "1~2 週", "2~3 個月"),
]
for i, (dim, v1, v2) in enumerate(rows):
    y = Inches(2.48) + Inches(i * 0.42)
    bg = CARD_BG if i % 2 == 0 else DARK_BLUE
    for j, val in enumerate([dim, v1, v2]):
        x = Inches(7.0) + Inches(j * 1.95)
        cell = add_rect(slide, x, y, Inches(1.85), Inches(0.40), bg)
        c = WHITE if j == 0 else (TEAL if j == 1 else ORANGE)
        b = True if j == 0 else False
        set_shape_text(cell, val, 8, c, b)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 6, TOTAL_SLIDES)

# ============================================================
# SLIDE 7: Agent 開發者 — 管控與軌跡（Harness 治理外框）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "05  Agent 開發者 — Harness 治理外框與管控軌跡", 24, WHITE, True)
add_text_box(slide, Inches(0.8), Inches(0.85), Inches(11), Inches(0.4),
             "模型負責「想」，Harness 負責「准不准做」── 四項強制功能 + 五層管控 + 8欄位稽核", 12, LIGHT_GRAY, False)

# Four mandatory functions (from 整合提報版 Slide 8)
add_text_box(slide, Inches(0.5), Inches(1.3), Inches(5), Inches(0.35), "Harness 四項強制功能（無法關閉、無法改寫）", 13, GOLD, True)
funcs = [
    ("1 工具白名單", "AI 能動用哪些工具，須先經審查登記；外部來源之 MCP 工具伺服器一律不採用，全數自建自審"),
    ("2 重大動作先問人", "對外發送、系統寫入等無法復原之動作，執行前一律停下等候人工核准 (HITL interrupt())"),
    ("3 用量與成本上限", "單次任務與各單位每日用量皆設上限，異常暴衝自動熔斷，避免失控支出"),
    ("4 全程留紀錄", "做了什麼、查了哪些資料、產出什麼，完整留存 (append-only)，事後可逐步重建與課責"),
]
for i, (title, desc) in enumerate(funcs):
    y = Inches(1.75) + Inches(i * 0.55)
    add_text_box(slide, Inches(0.5), y, Inches(0.35), Inches(0.25), title[:1], 11, GOLD, True)
    add_text_box(slide, Inches(0.85), y, Inches(1.5), Inches(0.25), title[2:], 11, ACCENT_BLUE, True)
    add_text_box(slide, Inches(2.4), y, Inches(4.0), Inches(0.45), desc, 9, LIGHT_GRAY, False)

# Right: L1-L5 Harness coverage (from 整合提報版 Slide 7)
add_shape(slide, Inches(6.8), Inches(1.3), Inches(6.2), Inches(3.1), CARD_BG, PURPLE, Pt(1.5))
add_text_box(slide, Inches(7.0), Inches(1.4), Inches(5), Inches(0.3), "Harness 貫穿 L1~L5 五層強制作為", 12, GOLD, True)
layers_harness = [
    ("L5 知識庫", "檢索標記不可信 · 依密級過濾 · DLP比對"),
    ("L4 技能/MCP", "簽章白名單 · Schema驗證 · 沙箱執行"),
    ("L3 Agent", "迴圈步數熔斷 · 部門隔離 · HITL閘門"),
    ("L2 DSLM", "語料核可清單 · 評測門檻 gating"),
    ("L1 基礎模型", "雜湊驗簽 · 系統提示凍結唯讀 · 參數上限"),
]
for i, (layer, desc) in enumerate(layers_harness):
    y = Inches(1.8) + Inches(i * 0.47)
    add_text_box(slide, Inches(7.0), y, Inches(1.5), Inches(0.25), layer, 9, ORANGE, True)
    add_text_box(slide, Inches(8.5), y, Inches(4.3), Inches(0.4), desc, 9, LIGHT_GRAY, False)

# Bottom: 8-field audit record
add_shape(slide, Inches(0.3), Inches(4.6), Inches(12.7), Inches(2.5), CARD_BG, GREEN, Pt(1))
add_text_box(slide, Inches(0.5), Inches(4.7), Inches(5), Inches(0.35), "8 欄位不可否認稽核紀錄（防竄改保存 ≥ 1 年）", 13, GOLD, True)
audit_fields = [
    ("① 請求識別碼", "Request ID / Trace ID — 全鏈路追溯"),
    ("② 身分/租戶", "User/Tenant ID — SPIFFE/SPIRE 工作負載憑證"),
    ("③ 模型版本", "Model Version + 組態 — 含 SBOM 紀錄"),
    ("④ 遮罩提示詞", "Masked Prompt — 脫敏後完整保留"),
    ("⑤ RAG 來源", "引用來源+DISA IL 分級標籤 — 命中紀錄"),
    ("⑥ 工具呼叫", "Tool Calls + 參數 + 返回值 — 完整捕獲"),
    ("⑦ HITL 核可", "Approver Signature — 審批者身分+時戳"),
    ("⑧ 處置結果", "Decision + Status — allow/redact/block/escalate"),
]
for i, (field, desc) in enumerate(audit_fields):
    col_x = Inches(0.5) + Inches((i % 4) * 3.15)
    row_y = Inches(5.15) + Inches((i // 4) * 0.65)
    add_text_box(slide, col_x, row_y, Inches(1.4), Inches(0.25), field, 10, GREEN, True)
    add_text_box(slide, col_x, row_y + Inches(0.25), Inches(3.0), Inches(0.3), desc, 9, LIGHT_GRAY, False)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 7, TOTAL_SLIDES)

# ============================================================
# SLIDE 8: Agent 使用者 — Portal 操作流程（對齊整合提報版 Slide 4）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "06  Agent 使用者 — Portal 到 Agent 的 11 步資料流", 24, WHITE, True)

# 11-step flow (aligned with 整合提報版 Slide 4)
steps_11 = [
    ("① 登入提需求", "瀏覽器/用戶端\n登入 Portal"),
    ("② 身分驗證", "SSO+MFA\n角色套用"),
    ("③ 組裝請求", "範本+語料選集\n+分級標記"),
    ("④ 身分識別", "Gateway\n部門識別"),
    ("⑤ 分級檢核", "注入防護\n密級比對"),
    ("⑥ 模型路由", "落點判定\n地端/雲端"),
    ("⑦ 推論執行", "工具呼叫\n向量檢索"),
    ("⑧ 輸出過濾", "DLP + 來源\n標註校驗"),
    ("⑨ 稽核落檔", "唯讀累加\n8欄位留存"),
    ("⑩ 回應呈現", "引用標示\n分級標籤"),
    ("⑪ 取得回應", "使用者\n下載/查閱"),
]
for i, (title, desc) in enumerate(steps_11):
    x = Inches(0.15) + Inches(i * 1.2)
    color = [ACCENT_BLUE, ACCENT_BLUE, TEAL, ORANGE, ORANGE, ORANGE, MID_BLUE, PURPLE, GREEN, TEAL, ACCENT_BLUE][i]
    card = add_shape(slide, x, Inches(1.2), Inches(1.08), Inches(1.5), CARD_BG, color, Pt(1))
    tf = card.text_frame; tf.word_wrap = True; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(8); p.font.color.rgb = color; p.font.bold = True; p.font.name = "Microsoft JhengHei"
    p2 = tf.add_paragraph(); p2.text = desc; p2.font.size = Pt(7); p2.font.color.rgb = LIGHT_GRAY; p2.font.name = "Microsoft JhengHei"; p2.alignment = PP_ALIGN.CENTER
    if i < 10:
        add_arrow_shape(slide, x + Inches(1.1), Inches(1.75), Inches(0.15), Inches(0.18), GOLD)

# HITL flow detail
add_shape(slide, Inches(0.3), Inches(3.0), Inches(12.7), Inches(2.0), CARD_BG, ORANGE, Pt(1.5))
add_text_box(slide, Inches(0.5), Inches(3.1), Inches(5), Inches(0.3), "HITL 人機審批流程（高衝擊動作強制中斷）", 13, GOLD, True)
hitl_steps = [
    ("Agent 提議動作", "模型產出工具\n呼叫意圖"),
    ("Harness 攔截", "interrupt_on\n掛起執行"),
    ("狀態快照保存", "Checkpointer\n保存斷點"),
    ("呈現審批面板", "Portal 顯示\n決策軌跡"),
    ("人工決策", "approve/edit\n/reject"),
    ("恢復執行", "Command\n(resume=...)"),
]
for i, (title, desc) in enumerate(hitl_steps):
    x = Inches(0.5) + Inches(i * 2.1)
    s = add_shape(slide, x, Inches(3.5), Inches(1.9), Inches(1.2), MID_BLUE, ORANGE, Pt(1))
    tf = s.text_frame; tf.word_wrap = True; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(9); p.font.color.rgb = ORANGE; p.font.bold = True; p.font.name = "Microsoft JhengHei"
    p2 = tf.add_paragraph(); p2.text = desc; p2.font.size = Pt(8); p2.font.color.rgb = LIGHT_GRAY; p2.font.name = "Microsoft JhengHei"; p2.alignment = PP_ALIGN.CENTER
    if i < 5:
        add_arrow_shape(slide, x + Inches(1.93), Inches(3.95), Inches(0.2), Inches(0.22), ORANGE)

# Key controls
add_shape(slide, Inches(0.3), Inches(5.3), Inches(12.7), Inches(1.9), CARD_BG, TEAL, Pt(1))
add_text_box(slide, Inches(0.5), Inches(5.4), Inches(5), Inches(0.3), "使用者端管控機制", 13, GOLD, True)
controls = [
    ("Portal 管控", "身分+部門綁定 · Agent/技能上架審查 · 用途+密級宣告 · 禁跨部門歷史 · 閒置逾期自動下架"),
    ("Gateway 管控", "OPA/Rego 三維裁決 · 金鑰集中 · 速率/併發/成本上限 · 全量留跡+歸屬計費 · fail-closed"),
    ("隔離原則", "會話 thread_id 隔離 · namespace 多租戶 · 部門 GPU 配額 · 跨密級存取告警 · 即時降級/鎖定"),
]
for i, (title, desc) in enumerate(controls):
    y = Inches(5.8) + Inches(i * 0.42)
    add_text_box(slide, Inches(0.6), y, Inches(1.5), Inches(0.25), "▶ " + title, 10, TEAL, True)
    add_text_box(slide, Inches(2.2), y, Inches(10.5), Inches(0.35), desc, 9, LIGHT_GRAY, False)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 8, TOTAL_SLIDES)

# ============================================================
# SLIDE 9: Agent 使用者 — 管控與軌跡
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "07  Agent 使用者 — 操作管控與軌跡紀錄", 24, WHITE, True)

# Left: Portal/Gateway dual control
add_shape(slide, Inches(0.3), Inches(1.1), Inches(6.2), Inches(6.0), CARD_BG, TEAL, Pt(1.5))
add_text_box(slide, Inches(0.5), Inches(1.2), Inches(5), Inches(0.4), "Portal + Gateway 雙管控模型", 14, GOLD, True)

dual_items = [
    ("Portal 側管控", ACCENT_BLUE, [
        "• 單一撤銷點：一鍵停用部門/模型/Agent",
        "• 上架審查：Agent+技能須經安全審查後掛載",
        "• 用途宣告：申請時宣告最高資料密級",
        "• 使用可視化：用量/成本/拒絕率儀表板",
        "• 生命週期收斂：閒置逾期自動下架",
    ]),
    ("Gateway 側強制", ORANGE, [
        "• G1 強制通道：全院 AI 僅經 Gateway",
        "• G2 落點路由：依分級決定地端/圍籬/外部",
        "• G3 流量成本：逾限降級或拒絕 (fail-closed)",
        "• G4 稽核歸屬：全量留存+部門歸屬計費",
        "• 金鑰集中：用戶端與 Agent 一律不持有",
    ]),
    ("RBAC + ABAC", PURPLE, [
        "• 角色 RBAC：管理員/開發者/使用者/審核官",
        "• 屬性 ABAC：單位/密級/時段/裝置+網段",
        "• 短時效 Token：mTLS 服務身分識別",
        "• 最小權限：動態策略引擎即時評估",
        "• 越權即鎖：異常行為即時降級封鎖",
    ]),
]
y_pos = Inches(1.7)
for section_title, color, items in dual_items:
    add_text_box(slide, Inches(0.5), y_pos, Inches(5.5), Inches(0.25), "◆ " + section_title, 11, color, True)
    y_pos += Inches(0.3)
    for item in items:
        add_text_box(slide, Inches(0.7), y_pos, Inches(5.5), Inches(0.22), item, 9, LIGHT_GRAY, False)
        y_pos += Inches(0.25)
    y_pos += Inches(0.12)

# Right: Audit trail
add_shape(slide, Inches(6.8), Inches(1.1), Inches(6.2), Inches(6.0), CARD_BG, GREEN, Pt(1.5))
add_text_box(slide, Inches(7.0), Inches(1.2), Inches(5), Inches(0.4), "不可否認軌跡與 fail-closed 原則", 14, GOLD, True)

audit_sections = [
    ("append-only 稽核體系", TEAL, [
        "• 僅可附加 JSONL 記錄格式",
        "• 提示+工具呼叫+結果完整留存",
        "• 雜湊鏈串接+異地備份",
        "• 保存期限 ≥ 1 年（法規要求）",
    ]),
    ("fail-closed 設計原則", RED, [
        "• 檢查器不可用或逾時→一律不通過",
        "• 請求未匹配核准政策→一律拒絕",
        "• 前一節點失效→不得使後續失守",
        "• HITL 逾時→預設拒絕",
    ]),
    ("處置矩陣（五種判定）", ORANGE, [
        "• 放行 (allow)：全部通過留存判定",
        "• 去識別 (redact)：遮蔽個資後放行",
        "• 改寫重試 (rewrite)：退回限次重生成",
        "• 阻斷 (block)：機密/金鑰/有害→終止告警",
        "• 轉人工 (escalate)：高衝擊→HITL 裁決",
    ]),
    ("SOC 介接", GREEN, [
        "• 即時告警：異常頻率/跨密級/越權",
        "• 斷路器 (Circuit Breaker) 自動熔斷",
        "• 事故復原：權杖撤銷+語料回收",
    ]),
]
y_pos = Inches(1.7)
for section_title, color, items in audit_sections:
    add_text_box(slide, Inches(7.0), y_pos, Inches(5.5), Inches(0.25), "◆ " + section_title, 11, color, True)
    y_pos += Inches(0.3)
    for item in items:
        add_text_box(slide, Inches(7.2), y_pos, Inches(5.5), Inches(0.22), item, 9, LIGHT_GRAY, False)
        y_pos += Inches(0.24)
    y_pos += Inches(0.08)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 9, TOTAL_SLIDES)

# ============================================================
# SLIDE 10: 資安防護（一）：單閘 Diode · DLP
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "08  資安防護（一）：單閘 Diode · DLP · 落點路由", 24, WHITE, True)

# Diode
add_shape(slide, Inches(0.3), Inches(1.1), Inches(4.0), Inches(5.6), CARD_BG, ORANGE, Pt(2))
add_text_box(slide, Inches(0.5), Inches(1.2), Inches(3.5), Inches(0.35), "🔒 單向閘道器 (Diode/CDS)", 13, GOLD, True)
diode_lines = [
    ("入向管制", TEAL, [
        "物理單向傳輸 · 深度封包檢測",
        "外網語料落地隔離區",
        "格式正規化：剔除巨集/可執行碼",
        "僅保留純文字/結構化資料",
        "人工抽驗→通過後匯入內網",
    ]),
    ("出向管制", RED, [
        "機敏資料絕對禁止離開內網",
        "禁止送往公有雲模型",
        "所有出向→DLP+PII遮罩檢核",
        "嚴禁任何形式繞道直連",
    ]),
]
y_pos = Inches(1.65)
for title, color, items in diode_lines:
    add_text_box(slide, Inches(0.5), y_pos, Inches(3.5), Inches(0.25), "▸ " + title, 10, color, True)
    y_pos += Inches(0.28)
    for item in items:
        add_text_box(slide, Inches(0.6), y_pos, Inches(3.4), Inches(0.22), item, 9, LIGHT_GRAY, False)
        y_pos += Inches(0.24)
    y_pos += Inches(0.08)

# DLP
add_shape(slide, Inches(4.5), Inches(1.1), Inches(4.0), Inches(5.6), CARD_BG, RED, Pt(2))
add_text_box(slide, Inches(4.7), Inches(1.2), Inches(3.5), Inches(0.35), "🛡️ DLP + PII 遮罩引擎", 13, GOLD, True)
dlp_lines = [
    ("DLP 檢核層", ORANGE, [
        "即時內容掃描引擎",
        "機敏關鍵字/正則/ML 分類器",
        "文件指紋 (Fingerprint)",
        "影像 OCR 文字提取檢核",
    ]),
    ("PII 遮罩", TEAL, [
        "身分證號/電話/地址→自動脫敏",
        "軍事座標/武器參數→等級遮罩",
        "人名/單位→去識別化",
        "金鑰/密級標記→即時阻斷",
    ]),
    ("處置機制", RED, [
        "即時阻斷+告警通報",
        "事件升級 SOC 監控中心",
        "違規紀錄完整留存",
    ]),
]
y_pos = Inches(1.65)
for title, color, items in dlp_lines:
    add_text_box(slide, Inches(4.7), y_pos, Inches(3.5), Inches(0.25), "▸ " + title, 10, color, True)
    y_pos += Inches(0.28)
    for item in items:
        add_text_box(slide, Inches(4.8), y_pos, Inches(3.4), Inches(0.22), item, 9, LIGHT_GRAY, False)
        y_pos += Inches(0.24)
    y_pos += Inches(0.08)

# Data routing
add_shape(slide, Inches(8.7), Inches(1.1), Inches(4.3), Inches(5.6), CARD_BG, PURPLE, Pt(2))
add_text_box(slide, Inches(8.9), Inches(1.2), Inches(3.8), Inches(0.35), "📍 三落點路由判定", 13, GOLD, True)
route_lines = [
    ("A 地端院內機房", ACCENT_BLUE, [
        "D0~D3 皆可存放",
        "D3 僅限實體隔離網域",
        "本院自建模型+AI助理",
        "不依賴任何外部服務",
    ]),
    ("B 圍籬內雲端 (VPC)", TEAL, [
        "D0/D1 可存放",
        "D2 須逐案核可",
        "CMEK/EKM 自管金鑰",
        "服務邊界限制資料外流",
    ]),
    ("C 圍籬外雲端", RED, [
        "不存放本院任何資料",
        "僅供 D0+去識別化 D1",
        "經閘道呼叫+逐案核准",
        "契約載明零資料保留",
    ]),
]
y_pos = Inches(1.65)
for title, color, items in route_lines:
    add_text_box(slide, Inches(8.9), y_pos, Inches(3.8), Inches(0.25), "▸ " + title, 10, color, True)
    y_pos += Inches(0.28)
    for item in items:
        add_text_box(slide, Inches(9.0), y_pos, Inches(3.7), Inches(0.22), item, 9, LIGHT_GRAY, False)
        y_pos += Inches(0.24)
    y_pos += Inches(0.08)

add_text_box(slide, Inches(0.5), Inches(6.8), Inches(12.3), Inches(0.5),
             "判定原則：以「資料被送到哪一區」認定落點，而非 AI 服務位於何處 │ 就近且從嚴：地端能達成者不逕用雲端",
             9, GOLD, True, PP_ALIGN.CENTER)
add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 10, TOTAL_SLIDES)

# ============================================================
# SLIDE 11: 資安防護（二）：資料治理 · Guardrail（對齊整合提報版 Slide 11-12）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "09  資安防護（二）：資料治理管線 · Guardrail 五類佈設", 24, WHITE, True)

# Data governance pipeline (from 整合提報版 Slide 15)
add_shape(slide, Inches(0.3), Inches(1.1), Inches(6.2), Inches(5.9), CARD_BG, TEAL, Pt(2))
add_text_box(slide, Inches(0.5), Inches(1.2), Inches(5), Inches(0.35), "📊 Agentic 資料安全管控五步路徑", 13, GOLD, True)

dg_steps = [
    ("01 資料盤點與分級", ACCENT_BLUE, "依業務盤點資產 · 四級(公開/內部/敏感/機密)\n明訂禁入清單與跨域流通規則"),
    ("02 標籤綁定與繼承", TEAL, "分級寫入檔案屬性+資料庫欄位\n切塊/向量繼承來源標籤 · Agent 不得自行降級"),
    ("03 政策即程式碼", ORANGE, "OPA/Rego 轉為可執行政策\n身分×分級×用途 三維授權 · 版本化管理"),
    ("04 執行期強制點", PURPLE, "先授權後召回：檢索前過濾，非事後遮蔽\n工具呼叫+輸出雙向檢查 · 高衝擊 HITL 攔截"),
    ("05 稽核與復原", GREEN, "全鏈路留跡(提示/檢索/工具/輸出)\n越級存取即時告警+自動封鎖 · 權杖撤銷+語料回收"),
]
for i, (step, color, desc) in enumerate(dg_steps):
    y = Inches(1.65) + Inches(i * 0.92)
    badge = add_shape(slide, Inches(0.5), y, Inches(2.2), Inches(0.35), color)
    set_shape_text(badge, step, 9, WHITE, True)
    add_multiline_text(slide, Inches(2.85), y, Inches(3.3), Inches(0.8),
                       [(line, LIGHT_GRAY, False, 9) for line in desc.split("\n")], 9, LIGHT_GRAY, False, 1.4)
    if i < 4:
        add_down_arrow(slide, Inches(1.4), y + Inches(0.4), Inches(0.25), Inches(0.35), color)

# Guardrail 5 types (from 整合提報版 Slide 11)
add_shape(slide, Inches(6.8), Inches(1.1), Inches(6.2), Inches(5.9), CARD_BG, PURPLE, Pt(2))
add_text_box(slide, Inches(7.0), Inches(1.2), Inches(5), Inches(0.35), "🚧 Guardrail 五類佈設點（Gateway + Harness 雙點）", 12, GOLD, True)

gr_types = [
    ("輸入護欄", TEAL, "身分+部門檢核 · Prompt Injection 偵測\n輸入長度+外部內容來源標記"),
    ("檢索護欄", ACCENT_BLUE, "依密級過濾索引 · 剝除文件夾帶指令語意\n標記檢索內容為不可信"),
    ("工具/動作護欄", ORANGE, "工具白名單 · 參數 Schema 驗證\n副作用與資源範圍檢查"),
    ("輸出護欄", RED, "結構驗證 · 有害內容過濾 · DLP 掃描\n引用比對 · 去指令化防二次注入"),
    ("行為/迴圈護欄", PURPLE, "步數/代幣/成本上限 · 重複動作偵測\n逾限自動熔斷+收斂工作階段"),
]
for i, (title, color, desc) in enumerate(gr_types):
    y = Inches(1.7) + Inches(i * 0.92)
    badge = add_shape(slide, Inches(7.0), y, Inches(1.6), Inches(0.35), color)
    set_shape_text(badge, title, 10, WHITE, True)
    add_multiline_text(slide, Inches(8.75), y, Inches(4.0), Inches(0.8),
                       [(line, LIGHT_GRAY, False, 9) for line in desc.split("\n")], 9, LIGHT_GRAY, False, 1.4)

# Shared bottom
add_rect(slide, Inches(0.3), Inches(6.3), Inches(12.7), Inches(0.9), CARD_BG, GOLD, Pt(1))
add_text_box(slide, Inches(0.5), Inches(6.35), Inches(12.3), Inches(0.75),
             "不可退讓之架構紅線 ▸ 隔離邊界不得因 Agent 專案而開孔 ▸ 敏感級以上資料不得送往雲端模型 "
             "▸ 高衝擊動作必經 HITL 人工核可 ▸ MCP 伺服器一律自建自審後始得上架\n"
             "設計原則：檢查器不可用或逾時，一律視為不通過 (fail-closed)；兩段式檢查：確定性規則先行，僅灰區交由模型評審",
             9, GOLD, True, PP_ALIGN.CENTER)
add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 11, TOTAL_SLIDES)

# ============================================================
# SLIDE 12: 資安防護（三）：AI 國際標準全覽（全面補充）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "10  AI 國際標準與法規合規全覽", 24, WHITE, True)
add_text_box(slide, Inches(0.8), Inches(0.85), Inches(11), Inches(0.4),
             "由治理制度向下貫穿至技術驗證與稽核佐證 — 四大支柱完整覆蓋", 12, LIGHT_GRAY, False)

# 4 pillars (aligned with 整合提報版 Slide 2)
pillars = [
    ("01 ISO 治理框架", ACCENT_BLUE, [
        ("ISO/IEC 42001", "AI 管理系統 (AIMS)"),
        ("ISO/IEC 27001·27002", "資安管理+控制措施"),
        ("ISO/IEC 27017·27018", "雲端安全+雲端個資"),
        ("ISO/IEC 27040", "儲存安全"),
        ("ISO/IEC 27701", "隱私資訊管理 (PIMS)"),
        ("ISO 17025", "實驗室能力稽核佐證"),
    ]),
    ("02 NIST 風險框架", TEAL, [
        ("AI RMF 1.0", "GOVERN/MAP/MEASURE/MANAGE"),
        ("AI 600-1", "GenAI 剖繪·12類風險對策"),
        ("SP 800-53 R5", "安全與隱私控制基線"),
        ("SP 800-218A", "GenAI 安全開發 (SSDF)"),
        ("SP 800-171 R3", "CUI 非機密受控保護"),
        ("CSF 2.0", "Govern 功能銜接 AI 治理"),
        ("AI 100-2", "對抗式 ML 攻擊分類法"),
        ("SP 800-88", "媒體清除銷毀標準"),
    ]),
    ("03 OWASP 技術風險", ORANGE, [
        ("LLM Top 10 (2025)", "注入·敏感洩漏·供應鏈"),
        ("Agentic AI 威脅", "記憶投毒·工具濫用·逸脫"),
        ("AISVS", "AI 應用安全驗證標準(分級)"),
        ("ML Security Top 10", "資料投毒·模型竊取·對抗"),
        ("AI Exchange", "威脅控制對映 ISO/NIST"),
        ("GenAI Red Teaming", "紅隊測試方法論指引"),
        ("LLM 治理檢核", "採購與部署盡職調查"),
    ]),
    ("04 威脅知識庫與法制", RED, [
        ("MITRE ATLAS", "AI 對抗戰術技術知識庫"),
        ("EU AI Act", "風險分級·GPAI義務·罰則"),
        ("CISA/NSA 指引", "AI 系統安全部署指引"),
        ("CNSSI 1253", "國安系統影響等級評定"),
        ("DISA CC SRG", "IL2~IL6 軍規分級"),
        ("資通安全管理法", "A級機關防護基準+通報"),
        ("個資法", "訓練資料蒐集利用界限"),
        ("AI 基本法(推動中)", "國內 AI 治理上位規範"),
    ]),
]
for i, (title, color, items) in enumerate(pillars):
    x = Inches(0.2) + Inches(i * 3.3)
    card = add_shape(slide, x, Inches(1.3), Inches(3.15), Inches(5.2), CARD_BG, color, Pt(1.5))
    badge = add_shape(slide, x + Inches(0.08), Inches(1.42), Inches(2.99), Inches(0.45), color)
    set_shape_text(badge, title, 11, WHITE, True)
    y_pos = Inches(2.0)
    for std_name, std_desc in items:
        add_text_box(slide, x + Inches(0.1), y_pos, Inches(1.45), Inches(0.22), std_name, 8, color, True)
        add_text_box(slide, x + Inches(1.55), y_pos, Inches(1.5), Inches(0.22), std_desc, 7, LIGHT_GRAY, False)
        y_pos += Inches(0.35)

# Bottom flow
add_rect(slide, Inches(0.2), Inches(6.65), Inches(12.9), Inches(0.6), CARD_BG, GOLD, Pt(1))
flow_labels = ["治理制度\nISO 42001 AIMS", "風險流程\nNIST AI RMF 1.0", "技術控制\nOWASP LLM Top 10\n+ AISVS", "威脅驗證\nMITRE ATLAS\n紅隊演練", "稽核佐證\nISO 17025/27001"]
flow_colors = [ACCENT_BLUE, TEAL, ORANGE, RED, GREEN]
for i, (label, fc) in enumerate(zip(flow_labels, flow_colors)):
    x = Inches(0.4) + Inches(i * 2.55)
    s = add_shape(slide, x, Inches(6.72), Inches(2.2), Inches(0.45), fc)
    set_shape_text(s, label.split("\n")[0], 8, WHITE, True)
    if i < 4:
        add_arrow_shape(slide, x + Inches(2.23), Inches(6.83), Inches(0.3), Inches(0.22), GOLD)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 12, TOTAL_SLIDES)

# ============================================================
# SLIDE 13: 資安防護（四）：SOC 監控與熔斷
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "11  資安防護（四）：SOC 監控 · 熔斷 · 紅隊演練", 24, WHITE, True)

soc_layers = [
    ("資料蒐集", ACCENT_BLUE, "Agent 執行日誌 · Gateway 存取紀錄 · DLP 告警 · 模型調用紀錄 · HITL 審批日誌 · VFS 操作 · 認證授權日誌"),
    ("SIEM 分析", TEAL, "日誌集中化 · 關聯分析引擎 · 機器學習異常偵測 · UEBA 使用者行為分析 · 威脅情資整合 · 時序異常識別"),
    ("偵測告警", ORANGE, "異常查詢頻率 · 跨密級存取試探 · 越權模型調用 · Prompt Injection · 資料外洩嘗試 · 異常 Token 消耗"),
    ("熔斷回應", RED, "自動熔斷斷路器 (Circuit Breaker)：異常暴衝→即時切斷 Agent · 緊急停止 Kill Switch · 即拋式環境銷毀 · 憑證凍結"),
    ("事後調查", PURPLE, "完整對話還原 · 決策鏈追溯 · 工具呼叫重放 · 根因分析 · 紅隊語料庫擴充 · 改善建議 · 規則更新"),
]
for i, (name, color, desc) in enumerate(soc_layers):
    y = Inches(1.2) + Inches(i * 0.98)
    badge = add_shape(slide, Inches(0.3), y, Inches(1.8), Inches(0.45), color)
    set_shape_text(badge, name, 11, WHITE, True)
    desc_box = add_shape(slide, Inches(2.3), y, Inches(10.7), Inches(0.75), CARD_BG, color, Pt(1))
    tf = desc_box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = desc; p.font.size = Pt(9); p.font.color.rgb = LIGHT_GRAY; p.font.name = "Microsoft JhengHei"
    if i < 4:
        add_down_arrow(slide, Inches(1.0), y + Inches(0.5), Inches(0.22), Inches(0.35), color)

# Red team
add_rect(slide, Inches(0.3), Inches(6.2), Inches(12.7), Inches(0.95), CARD_BG, RED, Pt(1.5))
add_text_box(slide, Inches(0.5), Inches(6.25), Inches(12.3), Inches(0.8),
             "🔴 季度 Red Team AI 紅隊攻防演練 ▸ 越獄提示攻擊 ▸ 間接提示注入 ▸ 工具越權測試 ▸ 跨密級存取試探 ▸ 記憶投毒 ▸ 沙箱逃逸\n"
             "量測指標：機密洩漏率 · 越獄成功率 · 誤攔率(FPR) · 引用覆蓋率 · 延遲預算 │ 每次更版以固定測試集重跑比對，每月抽樣人工複核",
             9, LIGHT_GRAY, False, PP_ALIGN.LEFT)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 13, TOTAL_SLIDES)

# ============================================================
# SLIDE 14: 端到端全景流程
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "AI 資安防護端到端工作流程", 26, WHITE, True)

e2e_steps = [
    ("外網\n資料源", RGBColor(0x7F, 0x8C, 0x8D)),
    ("Diode\n單向閘道", ORANGE),
    ("落地\n隔離區", RED),
    ("惡意掃描\n格式正規", TEAL),
    ("DLP\nPII遮罩", PURPLE),
    ("資料治理\n五步管線", ACCENT_BLUE),
    ("Guardrail\n五類佈設", GREEN),
    ("Agent\n執行層", LIGHT_BLUE),
    ("Guardrail\n輸出護欄", GREEN),
    ("DLP\n出向檢核", PURPLE),
    ("使用者\n交付", RGBColor(0x7F, 0x8C, 0x8D)),
]
for i, (label, color) in enumerate(e2e_steps):
    x = Inches(0.2) + Inches(i * 1.2)
    shape = add_shape(slide, x, Inches(1.2), Inches(1.05), Inches(0.8), color)
    set_shape_text(shape, label, 8, WHITE, True)
    if i < 10:
        add_arrow_shape(slide, x + Inches(1.08), Inches(1.45), Inches(0.18), Inches(0.22), GOLD)

add_rect(slide, Inches(0.2), Inches(2.2), Inches(12.9), Inches(0.4), CARD_BG, RED, Pt(1.5))
add_text_box(slide, Inches(0.5), Inches(2.23), Inches(12.3), Inches(0.35),
             "🔴 SOC 7×24 持續監控 — SIEM · 異常偵測 · 斷路器熔斷 · Kill Switch · 紅隊演練 · 稽核留痕", 10, RED, True, PP_ALIGN.CENTER)

add_rect(slide, Inches(0.2), Inches(2.7), Inches(12.9), Inches(0.4), CARD_BG, GOLD, Pt(1))
add_text_box(slide, Inches(0.5), Inches(2.73), Inches(12.3), Inches(0.35),
             "📋 合規：ISO 42001 · NIST AI RMF 1.0 · AI 600-1 · SP 800-53 · OWASP AISVS · LLM Top 10 · MITRE ATLAS · DISA IL · EU AI Act",
             9, GOLD, True, PP_ALIGN.CENTER)

add_text_box(slide, Inches(0.5), Inches(3.4), Inches(5), Inches(0.35), "完整資安防護檢核清單", 15, WHITE, True)
cl_left = [
    "☑ 單向閘道器 (Diode/CDS) — 物理單向傳輸+DPI",
    "☑ 入向惡意掃描 — 病毒/隱碼/注入偵測+格式正規化",
    "☑ DLP 即時掃描 — 機敏關鍵字/正則/ML/指紋/OCR",
    "☑ PII 遮罩 — 身分證/座標/武器參數/金鑰自動脫敏",
    "☑ 資料五步管線 — 盤點→標籤繼承→政策碼→強制→稽核",
    "☑ OPA/Rego 三維授權 — 身分×分級×用途",
    "☑ Guardrail 五類佈設 — 輸入/檢索/動作/輸出/行為",
    "☑ 處置矩陣 — allow/redact/rewrite/block/escalate",
]
cl_right = [
    "☑ 四級沙箱隔離 — L0語言→L1程序→L2容器→L3微VM",
    "☑ HITL interrupt() — 高衝擊動作強制人工核可",
    "☑ fail-closed 原則 — 檢查器逾時一律不通過",
    "☑ Harness 四項強制 — 白名單/先問人/上限/留紀錄",
    "☑ 三落點路由 — 地端/圍籬內/圍籬外依分級判定",
    "☑ 8 欄位 append-only 稽核 — 防竄改保存 ≥ 1 年",
    "☑ SOC 斷路器 — 異常暴衝自動熔斷+Kill Switch",
    "☑ 紅隊演練 — 季度攻防+固定測試集比對+人工複核",
]
for i, item in enumerate(cl_left):
    add_text_box(slide, Inches(0.5), Inches(3.85) + Inches(i * 0.38), Inches(6), Inches(0.3), item, 9, LIGHT_GRAY, False)
for i, item in enumerate(cl_right):
    add_text_box(slide, Inches(6.8), Inches(3.85) + Inches(i * 0.38), Inches(6), Inches(0.3), item, 9, LIGHT_GRAY, False)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 14, TOTAL_SLIDES)

# ============================================================
# SLIDE 15: Roadmap（對齊整合提報版 Slide 18）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "12  AI 服務鏈 Roadmap — 四階段建置期程", 24, WHITE, True)

phases = [
    ("Q1 基礎防禦網", ACCENT_BLUE, [
        "▸ Control Gateway 最小可行版上線",
        "▸ 資料分級 D0~D3 定案+機器判讀",
        "▸ 稽核紀錄格式與留存政策發布",
        "▸ 單向 Diode 閘道部署規劃",
        "▸ 基礎身分認證 (SSO+MFA)",
        "▸ Harness 基礎映像準備",
    ]),
    ("Q2 單一部門先導", TEAL, [
        "▸ AI Portal 首版：模型+工具目錄",
        "▸ 先導單位隔離驗證",
        "▸ 以租代購契約資安條款範本",
        "▸ 首批行政庶務 Agent PoC",
        "▸ 開箱即用 Harness 部署驗證",
        "▸ 可觀測性平台建置",
    ]),
    ("Q3 AI 助理串聯", ORANGE, [
        "▸ 跨所 AI 助理協作上線",
        "▸ 跨單位調用核准流程系統化",
        "▸ 雲端圍籬+閘道事件回送整合",
        "▸ 客製化 Harness 框架建立",
        "▸ DeepAgents 武器系統 Agent 試點",
        "▸ SOC SIEM 全功能上線",
    ]),
    ("Q4 全院合規上線", GREEN, [
        "▸ 全院單位完成接取",
        "▸ 關閉舊有直連端點",
        "▸ ISO 42001 內部稽核+管審",
        "▸ 季度紅隊演練啟動",
        "▸ 全規範交叉合規矩陣",
        "▸ 影子 AI 完全收斂",
    ]),
]
for i, (title, color, items) in enumerate(phases):
    x = Inches(0.3) + Inches(i * 3.25)
    card = add_shape(slide, x, Inches(1.2), Inches(3.05), Inches(5.2), CARD_BG, color, Pt(2))
    badge = add_shape(slide, x + Inches(0.08), Inches(1.33), Inches(2.89), Inches(0.5), color)
    set_shape_text(badge, title, 13, WHITE, True)
    add_multiline_text(slide, x + Inches(0.12), Inches(2.0), Inches(2.8), Inches(4.0),
                       [(item, LIGHT_GRAY if not item.startswith("▸") else color, item.startswith("▸"), 10) for item in items],
                       10, LIGHT_GRAY, False, 1.5)
    if i < 3:
        add_arrow_shape(slide, x + Inches(3.08), Inches(3.5), Inches(0.25), Inches(0.3), color)

# Annual KPI
add_rect(slide, Inches(0.3), Inches(6.55), Inches(12.7), Inches(0.7), CARD_BG, GOLD, Pt(1))
add_text_box(slide, Inches(0.5), Inches(6.6), Inches(12.3), Inches(0.55),
             "年度成效衡量 ▸ 100% 全院 AI 呼叫經由閘道，無直連端點 ▸ 0 件 D2/D3 資料出境事件 ▸ 各單位重複採購與影子 AI 完全收斂",
             11, GOLD, True, PP_ALIGN.CENTER)
add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 15, TOTAL_SLIDES)

# ============================================================
# SLIDE 16: 效益評估與 KPI
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "13  效益評估與 KPI 指標", 24, WHITE, True)

kpi_data = [
    ("閘道覆蓋率", "100%", GREEN, ["全院 AI 呼叫經 Gateway", "無直連端點殘留", "影子 AI 完全收斂", "旁路直連封鎖驗證"]),
    ("資料出境事件", "0 件", TEAL, ["D2/D3 資料零出境", "DLP+PII 全攔截", "越權操作 100% 阻斷", "機密洩漏率趨零"]),
    ("合規覆蓋率", "100%", ACCENT_BLUE, ["ISO 42001 全項覆蓋", "NIST AI RMF 完整對標", "OWASP AISVS 分級通過", "MITRE ATLAS 紅隊驗證"]),
    ("作業效率提升", "↑ 60%", GOLD, ["公文撰擬 8hr→1hr", "報表產製 4hr→30min", "知識檢索效率 ×10", "人力釋出再部署"]),
]
for i, (title, metric, color, items) in enumerate(kpi_data):
    x = Inches(0.3) + Inches(i * 3.25)
    card = add_shape(slide, x, Inches(1.1), Inches(3.05), Inches(2.5), CARD_BG, color, Pt(1.5))
    add_text_box(slide, x + Inches(0.1), Inches(1.2), Inches(2.85), Inches(0.3), title, 13, WHITE, True, PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.1), Inches(1.55), Inches(2.85), Inches(0.55), metric, 28, color, True, PP_ALIGN.CENTER)
    add_multiline_text(slide, x + Inches(0.1), Inches(2.2), Inches(2.85), Inches(1.0),
                       [(item, LIGHT_GRAY, False, 9) for item in items], 9, LIGHT_GRAY, False, 1.5)

# Risk matrix
add_shape(slide, Inches(0.3), Inches(3.85), Inches(12.7), Inches(3.3), CARD_BG, PURPLE, Pt(1))
add_text_box(slide, Inches(0.5), Inches(3.95), Inches(5), Inches(0.35), "風險降低評估矩陣", 14, GOLD, True)

risk_headers = ["風險類別", "防護前", "防護後", "降低幅度", "主要控制措施"]
for i, h in enumerate(risk_headers):
    x = Inches(0.5) + Inches(i * 2.45)
    hdr = add_shape(slide, x, Inches(4.4), Inches(2.35), Inches(0.38), PURPLE)
    set_shape_text(hdr, h, 10, WHITE, True)

risk_rows = [
    ("資料外洩", "極高", "低", "↓ 90%", "Diode+DLP+PII+三落點路由"),
    ("Prompt Injection", "高", "極低", "↓ 95%", "Guardrail+格式正規+fail-closed"),
    ("越權操作", "高", "無", "↓ 100%", "Deny-First+HITL+四級沙箱+白名單"),
    ("模型幻覺", "中", "低", "↓ 70%", "事實驗證+引用比對+輸出護欄"),
    ("供應鏈攻擊", "中", "低", "↓ 80%", "MCP自建自審+SBOM+簽章驗證"),
    ("影子 AI", "高", "無", "↓ 100%", "Gateway強制通道+直連封鎖+Portal管理"),
]
for i, (risk, before, after, reduction, measure) in enumerate(risk_rows):
    y = Inches(4.83) + Inches(i * 0.40)
    bg = CARD_BG if i % 2 == 0 else DARK_BLUE
    colors_r = [WHITE, RED, GREEN, GOLD, LIGHT_GRAY]
    bolds_r = [True, True, True, True, False]
    for j, (val, c, b) in enumerate(zip([risk, before, after, reduction, measure], colors_r, bolds_r)):
        x = Inches(0.5) + Inches(j * 2.45)
        cell = add_rect(slide, x, y, Inches(2.35), Inches(0.38), bg)
        set_shape_text(cell, val, 9, c, b)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 16, TOTAL_SLIDES)

# ============================================================
# SLIDE 17: 結論與下一步行動
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOLD)
add_text_box(slide, Inches(0.8), Inches(0.3), Inches(10), Inches(0.6), "14  結論與下一步行動", 26, WHITE, True)

# Achievements
add_shape(slide, Inches(0.3), Inches(1.1), Inches(6.2), Inches(3.8), CARD_BG, TEAL, Pt(1.5))
add_text_box(slide, Inches(0.5), Inches(1.2), Inches(5), Inches(0.35), "🏆 關鍵研究成果", 15, GOLD, True)
achievements = [
    "✅ 五大原則驅動 AI 服務鏈完整架構設計",
    "✅ Harness 治理外框四項強制功能定義",
    "✅ 開箱即用 vs 客製化 Harness 雙軌開發框架",
    "✅ Portal/Gateway 11 步資料流全景設計",
    "✅ 端到端資安防護 16 項完整檢核清單",
    "✅ AI 國際標準四大支柱全覆蓋矩陣",
    "✅ Guardrail 五類佈設 + 處置矩陣 + fail-closed",
    "✅ SOC 熔斷 + 紅隊演練 + 8 欄位稽核體系",
    "✅ 三落點路由 + 資料五步管線 + OPA/Rego",
    "✅ 四階段 Roadmap + 年度 KPI 衡量指標",
]
for i, item in enumerate(achievements):
    add_text_box(slide, Inches(0.5), Inches(1.65) + Inches(i * 0.3), Inches(5.8), Inches(0.25), item, 9, LIGHT_GRAY, False)

# Next steps
add_shape(slide, Inches(6.8), Inches(1.1), Inches(6.2), Inches(3.8), CARD_BG, ORANGE, Pt(1.5))
add_text_box(slide, Inches(7.0), Inches(1.2), Inches(5), Inches(0.35), "🎯 下一步行動建議", 15, GOLD, True)
next_steps = [
    ("即刻啟動", GREEN, ["• Control Gateway 最小可行版建置", "• 資料分級 D0~D3 定案+機器判讀標籤"]),
    ("30 天內", TEAL, ["• 首批行政庶務 Agent PoC (開箱即用 Harness)", "• 稽核紀錄格式與留存政策發布"]),
    ("90 天內", ACCENT_BLUE, ["• AI Portal 首版上線 + 先導單位隔離驗證", "• Guardrail 五類佈設 + SOC 日誌串接"]),
    ("180 天內", ORANGE, ["• DeepAgents 客製化 Harness 開發啟動", "• ISO 42001 內部稽核 + 管理審查"]),
]
y_pos = Inches(1.65)
for title, color, items in next_steps:
    add_text_box(slide, Inches(7.0), y_pos, Inches(2), Inches(0.22), "▸ " + title, 10, color, True)
    y_pos += Inches(0.25)
    for item in items:
        add_text_box(slide, Inches(7.2), y_pos, Inches(5.5), Inches(0.2), item, 9, LIGHT_GRAY, False)
        y_pos += Inches(0.22)
    y_pos += Inches(0.06)

# Five principles final
add_shape(slide, Inches(0.3), Inches(5.15), Inches(12.7), Inches(2.1), CARD_BG, GOLD, Pt(2))
add_text_box(slide, Inches(0.5), Inches(5.25), Inches(12), Inches(0.35), "五大原則貫穿 AI 服務鏈全生命週期", 16, GOLD, True, PP_ALIGN.CENTER)
for i, (p, c) in enumerate(zip(principles, badge_colors)):
    s = add_shape(slide, Inches(0.8) + Inches(i * 2.5), Inches(5.7), Inches(2.2), Inches(0.55), c)
    set_shape_text(s, p, 15, WHITE, True)
add_text_box(slide, Inches(0.5), Inches(6.4), Inches(12.3), Inches(0.7),
             "本報告基於 Claude Opus 4.6 (Thinking) 深度研析，整合「AI Agent Harness 整合提報版 20260906」完整內容\n"
             "涵蓋：Harness 治理外框 · Portal/Gateway 雙管控 · Guardrail 五類佈設 · 四大國際標準支柱 · SOC 熔斷 · 四階段 Roadmap",
             9, LIGHT_GRAY, False, PP_ALIGN.CENTER)

add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), GOLD)
add_slide_number(slide, 17, TOTAL_SLIDES)

# ─── Save ───
output_path = r"d:\JavaDO\總結報告_v2.pptx"
prs.save(output_path)
print(f"✅ 總結報告 v2 已成功生成：{output_path}")
print(f"   共 {TOTAL_SLIDES} 頁投影片")
print("   修正內容：")
print("   1. 開箱即用 Harness = Claude Code / OpenAI Codex / Antigravity 2.0")
print("   2. 客製化 Harness = LangChain DeepAgents")
print("   3. AI 國際標準全面補充：NIST RMF 1.0/AI 600-1/SP 800-53/OWASP AISVS/MITRE ATLAS 等")
print("   4. 對齊整合提報版 20260906：Harness治理外框/Portal-Gateway/Guardrail五類/fail-closed/三落點/紅隊演練")
