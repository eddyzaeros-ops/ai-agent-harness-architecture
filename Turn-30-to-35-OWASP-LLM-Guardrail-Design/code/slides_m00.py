"""
Module 00 Slides (01 ~ 06): Overview & Cross-Layer Defense-in-Depth Matrix
"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from slides_common import (
    apply_warm_background,
    add_header,
    add_footer,
    add_bullet_card,
    add_custom_table,
    TOTAL_SLIDES,
    ALL_TABLES,
    BG_CARD,
    BORDER_CARD,
    BORDER_DIVIDER,
    BRAND_SLATE,
    BRAND_SAGE,
    BRAND_TERRA,
    BRAND_OCHRE,
    TEXT_HEADLINE,
    TEXT_BODY,
    TEXT_MUTED,
    FONT_HEADING,
    FONT_BODY
)


def build_slide_01_cover(prs):
    """Slide 01: Warm Editorial Minimalism Cover."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_warm_background(slide)
    
    cat_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(4.5), Inches(0.36))
    cat_box.fill.solid()
    cat_box.fill.fore_color.rgb = RGBColor(0xE8, 0xE2, 0xD5)
    cat_box.line.color.rgb = BORDER_CARD
    cat_box.line.width = Pt(0.75)
    tf_c = cat_box.text_frame
    tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
    p_c = tf_c.paragraphs[0]
    p_c.alignment = PP_ALIGN.CENTER
    r_c = p_c.add_run()
    r_c.text = "✦  EXECUTIVE ARCHITECTURE SPECIFICATION  ·  2026"
    r_c.font.name = FONT_HEADING
    r_c.font.size = Pt(10)
    r_c.font.bold = True
    r_c.font.color.rgb = BRAND_SAGE
    
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.733), Inches(1.6))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = "OWASP AI 護欄體系架構設計規範全景"
    r_t.font.name = FONT_HEADING
    r_t.font.size = Pt(32)
    r_t.font.bold = True
    r_t.font.color.rgb = TEXT_HEADLINE
    
    p_sub = tf_t.add_paragraph()
    p_sub.space_before = Pt(8)
    r_sub = p_sub.add_run()
    r_sub.text = "LLM 基礎模型 · Agentic 代理人行為 · Skills 技能沙盒 · RAG 檢索增強 · MITRE ATLAS 攻擊鍊全維對齊"
    r_sub.font.name = FONT_BODY
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = BRAND_TERRA
    
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.7), Inches(11.733), Inches(0.015))
    div.fill.solid()
    div.fill.fore_color.rgb = BORDER_DIVIDER
    div.line.fill.background()
    
    meta_items = [
        ("涵蓋核心安全標準", "OWASP LLM / Agentic / Skills / RAG / ATLAS", "跨層縱深防禦體系標準"),
        ("核心防禦架構哲學", "零信任 (Zero-Trust) · 狀態約束 · 破除致命三要素", "無狀態過濾轉向有狀態監督"),
        ("合規與風險對齊", "ISO/IEC 42001 AIMS · NIST AI RMF · EU AI Act", "具備不可否認性審計追蹤"),
        ("發布與維護工作小組", "AI Harness Engineering Architecture Taskforce", "企業級生產環境防護指引")
    ]
    
    card_w = 2.78
    card_h = 2.3
    gap = 0.20
    for idx, (title, main_val, sub_val) in enumerate(meta_items):
        c_left = 0.8 + idx * (card_w + gap)
        c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(4.0), Inches(card_w), Inches(card_h))
        c.fill.solid()
        c.fill.fore_color.rgb = BG_CARD
        c.line.color.rgb = BORDER_CARD
        c.line.width = Pt(1)
        
        tb = slide.shapes.add_textbox(Inches(c_left + 0.16), Inches(4.16), Inches(card_w - 0.32), Inches(card_h - 0.32))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        r0 = p0.add_run()
        r0.text = f"0{idx+1}  {title}"
        r0.font.name = FONT_HEADING
        r0.font.size = Pt(8.5)
        r0.font.bold = True
        r0.font.color.rgb = BRAND_SAGE
        
        p1 = tf.add_paragraph()
        p1.space_before = Pt(8)
        r1 = p1.add_run()
        r1.text = main_val
        r1.font.name = FONT_HEADING
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_HEADLINE
        
        p2 = tf.add_paragraph()
        p2.space_before = Pt(6)
        r2 = p2.add_run()
        r2.text = sub_val
        r2.font.name = FONT_BODY
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = TEXT_MUTED

    add_footer(slide, 1, TOTAL_SLIDES)


def build_slide_02_overview_matrix(prs):
    """Slide 02: 五層縱深防禦全景與三層護欄協同流轉."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_warm_background(slide)
    add_header(slide, "五層縱深防禦架構全景與三層核心護欄協同流轉", "護欄體系總覽", "從傳統網路/API邊界邁向語意、行為、沙盒與檢索的多維防禦鏈路")
    
    col_w = 3.75
    gap = 0.24
    top = 1.68
    h = 5.25
    
    add_bullet_card(
        slide, left=0.8, top=top, width=col_w, height=h,
        title="縱深防禦五層全景 (L1~L5)",
        items=[
            "L1 API 閘道層：傳輸層 TLS 加密、IP 速率限制與 WAF/DDoS 抵禦。",
            "L2 語意閘門層：Prompt 注入檢驗、越獄模式過濾、PII 遮罩與 DLP。",
            "L3 代理人執行監督層：有限狀態機（FSM）約束、目標錨定與去話術 HITL。",
            "L4 技能工具沙盒層：微容器（gVisor/Firecracker）隔離、致命三要素切斷。",
            "L5 檢索增強隔離層：文檔入庫清洗、Pre-Retrieval 權限過濾、XML 標籤封裝。"
        ],
        tag="DEFENSE-IN-DEPTH",
        accent_color=BRAND_SAGE,
        header_bg=BRAND_SLATE
    )
    
    add_bullet_card(
        slide, left=0.8 + col_w + gap, top=top, width=col_w, height=h,
        title="三層核心護欄協同 (Three-Tier Synergy)",
        items=[
            "G1 LLM 語意邊界：充當對話出入閘口，攔截對抗性文本與提示詞外洩。",
            "G2 Agentic 行為狀態機：監督自主決策與規劃路徑，防止目標被惡意劫持。",
            "G3 Skills 工具沙盒：嚴格管控底層 API/代碼執行，落實最小權限隔離。",
            "動態跨層連動：G1 發現可疑注入即時降級 G2 授權權限，阻斷 G3 執行。",
            "記憶分艙保護：短期記憶迭代洗滌，長期記憶雜湊綁定，防止記憶庫投毒。"
        ],
        tag="CORE SYNERGY",
        accent_color=BRAND_TERRA,
        header_bg=BRAND_SLATE
    )
    
    add_bullet_card(
        slide, left=0.8 + (col_w + gap) * 2, top=top, width=col_w, height=h,
        title="全鏈路不可否認性審計與防護規格",
        items=[
            "8 大關鍵欄位：包含時間戳、Trace ID、客戶身分、輸入 SHA-256、判定結果等。",
            "HMAC 防竄改簽章：確保日誌具備完整法律證據力與事故鑑識價值。",
            "去話術 HITL (Human-in-the-Loop)：嚴禁代理人修辭粉飾，UI 呈現原始結構化參數。",
            "動態 JIT 短效憑證：拒絕長期靜態金鑰，每次工具調用前申請 60 秒權限。",
            "全鏈路追蹤：透過分散式 Trace ID 貫穿 Client、閘道、LLM 與沙盒調用。"
        ],
        tag="AUDIT & GOVERNANCE",
        accent_color=BRAND_OCHRE,
        header_bg=BRAND_SLATE
    )
    
    add_footer(slide, 2, TOTAL_SLIDES)


def build_slide_03_module_navigation_table(prs):
    """Slide 03: 五篇專題筆記核心導航與關聯 (Table 1)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_warm_background(slide)
    add_header(slide, "五大核心防禦模組防護重心與安全框架對齊矩陣", "模組導航矩陣", "全景梳理各防禦模組之防護重心、對齊安全標準與整體架構守備職責")
    
    t_data = ALL_TABLES[1]
    add_custom_table(
        slide, left=0.8, top=1.68, width=11.733, height=5.25,
        headers=t_data['headers'],
        rows_data=t_data['rows'],
        col_widths=[2.4, 2.3, 2.5, 4.5],
        col_alignments=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT],
        font_size=8.5
    )
    add_footer(slide, 3, TOTAL_SLIDES)


def build_slide_04_four_tier_synergy_flow(prs):
    """Slide 04: 四層協同防護架構全景流程圖 (Mermaid 1)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_warm_background(slide)
    add_header(slide, "四層協同防護矩陣架構全景流程圖 (Defense-in-Depth Architecture)", "護欄體系架構圖", "使用者請求穿越 L1~L5 各防線之資料流轉與即時攔截節點全景 (Mermaid 1 對齊)")
    
    col_w = 2.78
    gap = 0.20
    top = 1.68
    h = 4.20
    
    tiers = [
        ("01", "L1 單一安全閘道", "API Gateway & Ingress", BRAND_SLATE, [
            "單一出入口收斂：所有請求必經閘道，嚴禁繞道直連模型。",
            "傳輸層防禦：TLS 1.3 強制加密、IP 白名單與 Token 驗證。",
            "速率配額限制：防止泛洪 DoS 與自動化掃描探測。",
            "WAF / DDoS 攔截：初篩網路層威脅與惡意爬蟲請求。"
        ]),
        ("02", "G1: LLM 基礎模型護欄", "L2 語意出入閘門", BRAND_SAGE, [
            "入向檢驗：Llama-Guard / NeMo 偵測越獄與 Prompt 注入。",
            "語意隔離：XML 邊界標籤隔離 Context 與 Instructions。",
            "出向過濾：Microsoft Presidio 動態遮罩 PII 機敏實體。",
            "出向 DLP：防止系統提示詞洩漏與惡意 XSS 輸出。"
        ]),
        ("03", "G2: Agentic 行為護欄", "L3 執行監督器", BRAND_TERRA, [
            "目標錨定門：System Prompt 簽名唯讀，餘弦相似度防漂移。",
            "狀態機約束：有限狀態機（FSM）約束規劃路徑，嚴禁跳躍。",
            "JIT 動態憑證：調用工具前簽發 60 秒短效最小權限 Token。",
            "情境化去話術 HITL：高危操作展示原始參數，人類覆核。"
        ]),
        ("04", "G3: Skills 工具沙盒", "L4/L5 沙盒與檢索隔離", BRAND_OCHRE, [
            "破除致命三要素：機敏憑證、外網、不可信語料嚴格互斥。",
            "微沙盒隔離：gVisor 容器執行，唯讀臨時儲存，網路禁通配符。",
            "RAG 檢索前置過濾：強制 Pre-Filtering 權限下推至向量庫。",
            "不可信回傳封裝：<untrusted_tool_output> 阻斷二次注入。"
        ])
    ]
    
    for idx, (num, title, sub, color, bullets) in enumerate(tiers):
        c_left = 0.8 + idx * (col_w + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(top), Inches(col_w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_CARD
        card.line.width = Pt(1)
        
        header_h = 0.52
        hbar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(top), Inches(col_w), Inches(header_h))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = color
        hbar.line.fill.background()
        
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(c_left + 0.10), Inches(top + 0.12), Inches(0.28), Inches(0.28))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        pb = tf_b.paragraphs[0]
        pb.alignment = PP_ALIGN.CENTER
        rb = pb.add_run()
        rb.text = num
        rb.font.name = FONT_HEADING
        rb.font.size = Pt(8.5)
        rb.font.bold = True
        rb.font.color.rgb = color
        
        tb_h = slide.shapes.add_textbox(Inches(c_left + 0.44), Inches(top + 0.06), Inches(col_w - 0.52), Inches(header_h - 0.12))
        tf_h = tb_h.text_frame
        tf_h.word_wrap = True
        tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0
        p_h = tf_h.paragraphs[0]
        r_th = p_h.add_run()
        r_th.text = title
        r_th.font.name = FONT_HEADING
        r_th.font.size = Pt(9.5)
        r_th.font.bold = True
        r_th.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        p_sh = tf_h.add_paragraph()
        r_sh = p_sh.add_run()
        r_sh.text = sub
        r_sh.font.name = FONT_BODY
        r_sh.font.size = Pt(8.0)
        r_sh.font.color.rgb = RGBColor(0xD6, 0xCE, 0xBF)
        
        tb_c = slide.shapes.add_textbox(Inches(c_left + 0.14), Inches(top + header_h + 0.10), Inches(col_w - 0.28), Inches(h - header_h - 0.16))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        for idx_b, b in enumerate(bullets):
            p = tf_c.paragraphs[0] if idx_b == 0 else tf_c.add_paragraph()
            p.space_after = Pt(4)
            p.line_spacing = 1.15
            parts = b.split("：", 1)
            rh = p.add_run()
            rh.text = "• " + parts[0] + "："
            rh.font.name = FONT_HEADING
            rh.font.size = Pt(8.5)
            rh.font.bold = True
            rh.font.color.rgb = BRAND_SLATE
            
            rb = p.add_run()
            rb.text = parts[1]
            rb.font.name = FONT_BODY
            rb.font.size = Pt(8.0)
            rb.font.color.rgb = TEXT_BODY

    # Bottom Audit Box
    bot_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.92))
    bot_box.fill.solid()
    bot_box.fill.fore_color.rgb = BG_CARD
    bot_box.line.color.rgb = BORDER_CARD
    bot_box.line.width = Pt(1)
    
    tb_bot = slide.shapes.add_textbox(Inches(0.95), Inches(6.12), Inches(11.433), Inches(0.78))
    tf_bot = tb_bot.text_frame
    tf_bot.word_wrap = True
    tf_bot.margin_left = tf_bot.margin_top = tf_bot.margin_right = tf_bot.margin_bottom = 0
    p_b0 = tf_bot.paragraphs[0]
    rb0 = p_b0.add_run()
    rb0.text = "✦  底層基石：全鏈路不可否認性審計 (Non-Repudiation Audit Trail)  ➔  "
    rb0.font.name = FONT_HEADING
    rb0.font.size = Pt(9.5)
    rb0.font.bold = True
    rb0.font.color.rgb = BRAND_SAGE
    
    rb1 = p_b0.add_run()
    rb1.text = "所有四層處置判定結果、Prompt SHA-256 雜湊與沙盒執行參數，皆統一寫入 HMAC-SHA256 簽署之安全日誌，即時推送 SOC / SIEM 中心，具備完整法律證據力。"
    rb1.font.name = FONT_BODY
    rb1.font.size = Pt(8.5)
    rb1.font.color.rgb = TEXT_BODY

    add_footer(slide, 4, TOTAL_SLIDES)


def build_slide_05_defense_in_depth_scenario(prs):
    """Slide 05: 跨層威脅滲透與三階段縱深攔截機制 (Defense-in-Depth Scenario)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_warm_background(slide)
    add_header(slide, "跨層威脅滲透與三階段縱深攔截機制", "縱深防禦場景", "模擬真實多階段複合攻擊演進，解析各護欄層級之協同聯動與梯次阻斷機制")
    
    col_w = 3.75
    gap = 0.24
    top = 1.68
    stage_h = 3.42
    
    stages_data = [
        {
            "stage_num": "1",
            "title": "階段一：入向語意穿透 (Ingress)",
            "layer_tag": "L5 檢索閘門 ⨂ L2 語意閘門",
            "header_color": BRAND_SLATE,
            "badge_color": BRAND_SAGE,
            "threat_title": "威脅滲透樣態 (Attack Vector)",
            "threat_desc": "攻擊者透過工單系統上傳夾帶「間接提示詞注入 (Indirect Prompt Injection)」之 PDF 文檔，利用白底微小字型隱藏攻擊指令，誘騙模型讀取內部敏感憑證與資料庫金鑰。",
            "defense_title": "縱深攔截防線 (Defense Interception)",
            "defense_bullets": [
                "L5 檢索閘門：文檔解析沙盒清除微小字型，輕量 SLM 分類器檢測高對抗性指令特徵。",
                "L2 語意閘門：檢索回傳強制以 <untrusted_rag_context> 標籤封裝，聲明純事實屬性嚴禁執行。"
            ]
        },
        {
            "stage_num": "2",
            "title": "階段二：代理決策劫持 (Goal Hijack)",
            "layer_tag": "L3 代理人執行監督器",
            "header_color": BRAND_SLATE,
            "badge_color": BRAND_TERRA,
            "threat_title": "威脅滲透樣態 (Attack Vector)",
            "threat_desc": "若注入指令透過變形混淆突破語意閘門，進入 Agent 推理上下文，企圖竄改代理人執行目標（Goal Hijack），誘導其背離原始任務去呼叫內部資料庫備份工具。",
            "defense_title": "縱深攔截防線 (Defense Interception)",
            "defense_bullets": [
                "L3 目標錨定門：即時計算 Action 與根目標之語意餘弦相似度，偏離門檻即刻告警阻斷。",
                "L3 規劃重評門：高危操作強制觸發「情境化去話術 HITL」，由人工審批結構化參數。"
            ]
        },
        {
            "stage_num": "3",
            "title": "階段三：橫向移動與外傳 (Exfiltration)",
            "layer_tag": "L4 技能工具沙盒 (Micro-Sandbox)",
            "header_color": BRAND_SLATE,
            "badge_color": BRAND_OCHRE,
            "threat_title": "威脅滲透樣態 (Attack Vector)",
            "threat_desc": "假設攻擊者進一步試圖直接呼叫系統底層工具連線外部惡意伺服器（C2）將資料打包外傳，或嘗試寫入本地持久化排程任務實施提權。",
            "defense_title": "縱深攔截防線 (Defense Interception)",
            "defense_bullets": [
                "L4 技能微沙盒：嚴格執行「破除致命三要素」，處理機敏資料之工具強制 Disable-Net 禁外網。",
                "L4 系統調用審計：eBPF 實時捕獲異常連線嘗試，秒級觸發看門狗熔斷並阻斷進程。"
            ]
        }
    ]
    
    for idx, s in enumerate(stages_data):
        c_left = 0.8 + idx * (col_w + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(top), Inches(col_w), Inches(stage_h))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_CARD
        card.line.width = Pt(1)
        
        header_h = 0.48
        h_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(top), Inches(col_w), Inches(header_h))
        h_shape.fill.solid()
        h_shape.fill.fore_color.rgb = s["header_color"]
        h_shape.line.fill.background()
        
        badge_dia = 0.26
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(c_left + 0.12), Inches(top + 0.11), Inches(badge_dia), Inches(badge_dia))
        badge.fill.solid()
        badge.fill.fore_color.rgb = s["badge_color"]
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        pb = tf_b.paragraphs[0]
        pb.alignment = PP_ALIGN.CENTER
        rb = pb.add_run()
        rb.text = s["stage_num"]
        rb.font.name = FONT_HEADING
        rb.font.size = Pt(8.5)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        tb_h = slide.shapes.add_textbox(Inches(c_left + 0.44), Inches(top + 0.05), Inches(col_w - 0.52), Inches(header_h - 0.10))
        tf_h = tb_h.text_frame
        tf_h.word_wrap = True
        tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0
        p_h = tf_h.paragraphs[0]
        r_title = p_h.add_run()
        r_title.text = s["title"]
        r_title.font.name = FONT_HEADING
        r_title.font.size = Pt(9.5)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        p_sub = tf_h.add_paragraph()
        r_tag = p_sub.add_run()
        r_tag.text = f"守備層級：{s['layer_tag']}"
        r_tag.font.name = FONT_BODY
        r_tag.font.size = Pt(8)
        r_tag.font.color.rgb = RGBColor(0xD6, 0xCE, 0xBF)
        
        content_top = top + header_h + 0.10
        content_h = stage_h - header_h - 0.18
        tb_c = slide.shapes.add_textbox(Inches(c_left + 0.16), Inches(content_top), Inches(col_w - 0.32), Inches(content_h))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        
        p0 = tf_c.paragraphs[0]
        r_th = p0.add_run()
        r_th.text = f"▲  {s['threat_title']}"
        r_th.font.name = FONT_HEADING
        r_th.font.size = Pt(9)
        r_th.font.bold = True
        r_th.font.color.rgb = BRAND_TERRA
        
        p1 = tf_c.add_paragraph()
        p1.space_before = Pt(3)
        p1.line_spacing = 1.15
        r_tb = p1.add_run()
        r_tb.text = s["threat_desc"]
        r_tb.font.name = FONT_BODY
        r_tb.font.size = Pt(8.5)
        r_tb.font.color.rgb = TEXT_BODY
        
        p2 = tf_c.add_paragraph()
        p2.space_before = Pt(7)
        r_dh = p2.add_run()
        r_dh.text = f"🛡  {s['defense_title']}"
        r_dh.font.name = FONT_HEADING
        r_dh.font.size = Pt(9)
        r_dh.font.bold = True
        r_dh.font.color.rgb = BRAND_SAGE
        
        for b_text in s["defense_bullets"]:
            pb = tf_c.add_paragraph()
            pb.space_before = Pt(3)
            pb.line_spacing = 1.15
            parts = b_text.split("：", 1)
            r_head = pb.add_run()
            r_head.text = "• " + parts[0] + "："
            r_head.font.name = FONT_HEADING
            r_head.font.size = Pt(8.5)
            r_head.font.bold = True
            r_head.font.color.rgb = BRAND_SLATE
            
            r_body = pb.add_run()
            r_body.text = parts[1]
            r_body.font.name = FONT_BODY
            r_body.font.size = Pt(8)
            r_body.font.color.rgb = TEXT_BODY

    # 2. Bottom Container: 零信任 AI 四大落地原則
    bot_top = top + stage_h + 0.16
    bot_h = 1.72
    
    bot_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(bot_top), Inches(11.733), Inches(bot_h))
    bot_card.fill.solid()
    bot_card.fill.fore_color.rgb = BG_CARD
    bot_card.line.color.rgb = BORDER_CARD
    bot_card.line.width = Pt(1)
    
    tb_bt = slide.shapes.add_textbox(Inches(1.0), Inches(bot_top + 0.08), Inches(11.333), Inches(0.26))
    tf_bt = tb_bt.text_frame
    tf_bt.word_wrap = True
    tf_bt.margin_left = tf_bt.margin_top = tf_bt.margin_right = tf_bt.margin_bottom = 0
    p_bt = tf_bt.paragraphs[0]
    r_bt = p_bt.add_run()
    r_bt.text = "✦  零信任（Zero-Trust）AI 護欄架構四大落地實踐原則"
    r_bt.font.name = FONT_HEADING
    r_bt.font.size = Pt(10)
    r_bt.font.bold = True
    r_bt.font.color.rgb = BRAND_SLATE
    
    principles = [
        ("01", "永不信任輸入", "所有外部文字與檢索語料一律視為不可信，強制包覆 XML 結構化隔離標籤與分類器檢驗。", BRAND_SAGE),
        ("02", "行為狀態約束", "以有限狀態機（FSM）界定合法行動邊界，限制自主步數上限，杜絕失控執行。", BRAND_TERRA),
        ("03", "破除致命三要素", "機敏憑證、外網連線、不可信語料絕不同時賦予單一工具，微沙盒強制隔離。", BRAND_OCHRE),
        ("04", "全鏈路審計", "全流程記錄輸入 SHA-256、判定結果與 HMAC-SHA256 防竄改簽章，確保法律鑑識效力。", BRAND_SLATE)
    ]
    
    p_w = 2.68
    p_gap = 0.20
    for idx, (num, p_title, p_desc, p_color) in enumerate(principles):
        p_left = 1.0 + idx * (p_w + p_gap)
        
        bdia = 0.24
        p_badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(p_left), Inches(bot_top + 0.38), Inches(bdia), Inches(bdia))
        p_badge.fill.solid()
        p_badge.fill.fore_color.rgb = p_color
        p_badge.line.fill.background()
        tf_pb = p_badge.text_frame
        tf_pb.margin_left = tf_pb.margin_top = tf_pb.margin_right = tf_pb.margin_bottom = 0
        ppb = tf_pb.paragraphs[0]
        ppb.alignment = PP_ALIGN.CENTER
        rpb = ppb.add_run()
        rpb.text = num
        rpb.font.name = FONT_HEADING
        rpb.font.size = Pt(8)
        rpb.font.bold = True
        rpb.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        tb_pt = slide.shapes.add_textbox(Inches(p_left + 0.32), Inches(bot_top + 0.36), Inches(p_w - 0.34), Inches(0.26))
        tf_pt = tb_pt.text_frame
        tf_pt.word_wrap = True
        tf_pt.margin_left = tf_pt.margin_top = tf_pt.margin_right = tf_pt.margin_bottom = 0
        ppt = tf_pt.paragraphs[0]
        rpt = ppt.add_run()
        rpt.text = p_title
        rpt.font.name = FONT_HEADING
        rpt.font.size = Pt(9.5)
        rpt.font.bold = True
        rpt.font.color.rgb = BRAND_SLATE
        
        tb_pd = slide.shapes.add_textbox(Inches(p_left), Inches(bot_top + 0.68), Inches(p_w), Inches(bot_h - 0.74))
        tf_pd = tb_pd.text_frame
        tf_pd.word_wrap = True
        tf_pd.margin_left = tf_pd.margin_top = tf_pd.margin_right = tf_pd.margin_bottom = 0
        ppd = tf_pd.paragraphs[0]
        ppd.line_spacing = 1.15
        rpd = ppd.add_run()
        rpd.text = p_desc
        rpd.font.name = FONT_BODY
        rpd.font.size = Pt(8.5)
        rpd.font.color.rgb = TEXT_BODY
        
    add_footer(slide, 5, TOTAL_SLIDES)


def build_slide_06_attack_tree_interception(prs):
    """Slide 06: 階層式複合攻擊滲透路徑與各層攔截流程圖 (Pre 3)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_warm_background(slide)
    add_header(slide, "階層式複合攻擊滲透路徑與縱深攔截全景流程圖", "攻擊演進路徑", "對抗性 Prompt 突破、代理人目標漂移、底層代碼執行與知識庫投毒的多層次攔截體系 (Pre 3 對齊)")
    
    col_w = 5.75
    gap = 0.23
    top = 1.68
    h = 5.25
    
    add_bullet_card(
        slide, left=0.8, top=top, width=col_w, height=h,
        title="攻擊滲透路徑 (前段)：提示詞注入與意圖劫持",
        items=[
            "階段 1：提示詞注入攻擊 (Prompt Injection)：使用角色扮演越獄 (Jailbreak) 或 Base64 混淆字串，試圖繞過直接對話邊界。",
            "第一道防線攔截：L1 閘道與 L2 LLM 護欄即時解析編碼，Llama-Guard 標記惡意 Prompt，強制阻斷並記錄事件。",
            "階段 2：意圖劫持與非預期工具呼叫：注入指令突破語意層，誘騙 Agent 背離原始目標去調用備份工具打包資料庫。",
            "第二道防線攔截：L3 Agentic 護欄之語意漂移監控（Drift Detector）比對子任務與最初登錄的根目標，發現語意偏離。",
            "狀態機檢查器（State Machine）：判定「非授權 ACT 躍遷」，立即凍結 Agent 執行緒，強制觸發情境化 HITL。"
        ],
        tag="ATTACK STAGES 1 & 2",
        accent_color=BRAND_TERRA,
        header_bg=BRAND_SLATE
    )
    
    right_left = 0.8 + col_w + gap
    add_bullet_card(
        slide, left=right_left, top=top, width=col_w, height=h,
        title="攻擊滲透路徑 (後段)：代碼外傳與記憶投毒",
        items=[
            "階段 3：底層代碼執行與資料外傳 (Exfiltration)：攻擊者試圖調用 Bash/Python 工具發起 Raw Socket 連線將金鑰回傳至遠端 C2。",
            "第三道防線攔截：L4 Skills 護欄落實「破除致命三要素」原則，微容器沙盒 (gVisor) 強制執行 Network: none 禁外網。",
            "eBPF 核心監控：系統調用層級攔截 connect() 呼叫，毫秒級殺死受污染進程並告警。",
            "階段 4：污染長期記憶或內部知識庫 (Persistence)：將惡意注入指令存入對話歷史或知識庫，使後續所有使用者持續受害。",
            "第四道防線攔截：L5 RAG 護欄與 L3 記憶寫入閘門強制執行去腳本化清洗與 SHA-256 簽章驗證，阻斷投毒持久化。"
        ],
        tag="ATTACK STAGES 3 & 4",
        accent_color=BRAND_SAGE,
        header_bg=BRAND_SLATE
    )
    
    add_footer(slide, 6, TOTAL_SLIDES)
