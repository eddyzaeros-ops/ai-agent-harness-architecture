"""
OWASP AI Guardrail Design Master Presentation Generator
Adhering to Template 1 (Warm Editorial Minimalism)
Full Coverage of HTML Knowledge Base:
- 48 Comprehensive Slides
- 13 Full Comparison Tables (No truncation)
- 6 Architectural Sequence Flow Diagrams
- 5 Code / YAML / JSON / XML Technical Specifications
- 6 Defense-in-Depth Scenario & Kill Chain Workflow Diagrams
- 4 Zero-Trust Engineering Principles & Architecture Manifesto
"""

import sys
import os
import json
import html

# Add template1 module path
sys.path.append(r"C:\Users\calsa\.gemini\config\skills\pptx-template-1\scripts")

import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from template1_builder import (
    init_presentation,
    apply_warm_background,
    add_header,
    add_footer,
    add_bullet_card,
    BG_CANVAS,
    BG_CARD,
    BORDER_CARD,
    BORDER_DIVIDER,
    BRAND_SLATE,
    BRAND_SAGE,
    BRAND_TERRA,
    BRAND_OCHRE,
    TEXT_HEADLINE,
    TEXT_BODY,
    TEXT_MUTED,
    TABLE_HEADER_BG,
    TABLE_ROW_ALT,
    FONT_HEADING,
    FONT_BODY
)

TOTAL_SLIDES = 48

# Load extracted tables
with open(r'd:\JavaDO\OWASP html\all_tables_extracted.json', 'r', encoding='utf-8') as f:
    ALL_TABLES = {t['table_id']: t for t in json.load(f)}

# -----------------------------------------------------------------------------
# CORE TEMPLATE 1 REFINED BUILDERS
# -----------------------------------------------------------------------------

def add_custom_table(slide, left, top, width, height, headers, rows_data, col_widths=None, col_alignments=None, font_size=8.0, has_card_container=True):
    """Refined Table Component with customized cell padding and text styling."""
    if has_card_container:
        container = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        container.fill.solid()
        container.fill.fore_color.rgb = BG_CARD
        container.line.color.rgb = BORDER_CARD
        container.line.width = Pt(1)
        
        t_left = left + 0.12
        t_top = top + 0.10
        t_width = width - 0.24
        t_height = height - 0.20
    else:
        t_left = left
        t_top = top
        t_width = width
        t_height = height

    num_rows = len(rows_data) + 1
    num_cols = len(headers)
    table_shape = slide.shapes.add_table(num_rows, num_cols, Inches(t_left), Inches(t_top), Inches(t_width), Inches(t_height))
    table = table_shape.table
    
    if col_widths and len(col_widths) == num_cols:
        scale = t_width / sum(col_widths)
        for idx, w in enumerate(col_widths):
            table.columns[idx].width = Inches(w * scale)
            
    # Header Row
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_HEADER_BG
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Pt(5)
        cell.margin_right = Pt(5)
        cell.margin_top = Pt(4)
        cell.margin_bottom = Pt(4)
        
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        align = col_alignments[col_idx] if col_alignments and col_idx < len(col_alignments) else (PP_ALIGN.LEFT if col_idx == 0 else PP_ALIGN.CENTER)
        p.alignment = align
        p.font.name = FONT_HEADING
        p.font.size = Pt(font_size + 0.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_HEADLINE
        
    # Data Rows
    for row_idx, r_data in enumerate(rows_data):
        row_bg = TABLE_ROW_ALT if row_idx % 2 == 1 else BG_CARD
        for col_idx, val in enumerate(r_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_bg
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Pt(5)
            cell.margin_right = Pt(5)
            cell.margin_top = Pt(4)
            cell.margin_bottom = Pt(4)
            
            p = cell.text_frame.paragraphs[0]
            val_str = str(val).strip()
            
            lines = val_str.split("\n")
            p.text = lines[0]
            align = col_alignments[col_idx] if col_alignments and col_idx < len(col_alignments) else (PP_ALIGN.LEFT if col_idx == 0 else PP_ALIGN.CENTER)
            p.alignment = align
            p.font.name = FONT_BODY
            p.font.size = Pt(font_size)
            p.line_spacing = 1.12
            
            if col_idx == 0:
                p.font.bold = True
                p.font.color.rgb = TEXT_HEADLINE
            else:
                p.font.color.rgb = TEXT_BODY
                
            for extra_line in lines[1:]:
                p_next = cell.text_frame.add_paragraph()
                p_next.text = extra_line
                p_next.alignment = align
                p_next.font.name = FONT_BODY
                p_next.font.size = Pt(font_size)
                p_next.line_spacing = 1.12
                if col_idx == 0:
                    p_next.font.bold = True
                    p_next.font.color.rgb = TEXT_HEADLINE
                else:
                    p_next.font.color.rgb = TEXT_BODY


def add_quote_box(slide, left, top, width, height, en_quote, zh_translation, tag_text=None, accent_color=BRAND_TERRA):
    """Standard Centered Warm Editorial Quote Box."""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = BG_CARD
    box.line.color.rgb = BORDER_CARD
    box.line.width = Pt(1)
    
    tb = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.20), Inches(width - 0.6), Inches(height - 0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    if tag_text:
        r_tag = p0.add_run()
        r_tag.text = f"✦  {tag_text}\n\n"
        r_tag.font.name = FONT_HEADING
        r_tag.font.size = Pt(9.5)
        r_tag.font.bold = True
        r_tag.font.color.rgb = BRAND_SAGE
        
    r_en = p0.add_run()
    r_en.text = f"“ {en_quote} ”"
    r_en.font.name = FONT_HEADING
    r_en.font.size = Pt(13)
    r_en.font.bold = True
    r_en.font.color.rgb = accent_color
    
    p1 = tf.add_paragraph()
    p1.alignment = PP_ALIGN.CENTER
    p1.space_before = Pt(8)
    r_zh = p1.add_run()
    r_zh.text = zh_translation
    r_zh.font.name = FONT_BODY
    r_zh.font.size = Pt(9.5)
    r_zh.font.color.rgb = TEXT_MUTED


def add_code_box(slide, left, top, width, height, code_text, title=None, lang="YAML"):
    """Monospace Code / Specification Card."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xFB, 0xFA, 0xF7)
    card.line.color.rgb = BORDER_CARD
    card.line.width = Pt(1)
    
    h_top = 0.36
    if title:
        hbar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(h_top))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = TABLE_HEADER_BG
        hbar.line.fill.background()
        tf_hb = hbar.text_frame
        p_hb = tf_hb.paragraphs[0]
        r_hb = p_hb.add_run()
        r_hb.text = f"💻  {title}  [{lang.upper()}]"
        r_hb.font.name = "Consolas"
        r_hb.font.size = Pt(8.5)
        r_hb.font.bold = True
        r_hb.font.color.rgb = BRAND_SLATE
        c_top = top + h_top + 0.08
        c_h = height - h_top - 0.16
    else:
        c_top = top + 0.12
        c_h = height - 0.24

    tb = slide.shapes.add_textbox(Inches(left + 0.16), Inches(c_top), Inches(width - 0.32), Inches(c_h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = code_text.strip()
    p.font.name = "Consolas"
    p.font.size = Pt(8.0)
    p.font.color.rgb = BRAND_SLATE
    p.line_spacing = 1.15


def add_sequence_step_grid(slide, steps, top=1.68):
    """
    Renders 6 Sequence Flow Steps in a 2x3 Grid with Template 1 styling.
    Each step: (step_num, title, actor_direction, tag, desc_bullets, accent_color)
    """
    col_w = 3.75
    gap_x = 0.24
    row_h = 2.48
    gap_y = 0.20
    
    for idx, (num, title, actor_dir, tag, bullets, color) in enumerate(steps):
        r_idx = idx // 3
        c_idx = idx % 3
        
        c_left = 0.8 + c_idx * (col_w + gap_x)
        c_top = top + r_idx * (row_h + gap_y)
        
        # Outer Card
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(c_top), Inches(col_w), Inches(row_h))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_CARD
        card.line.width = Pt(1)
        
        # Header bar
        header_h = 0.44
        h_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(c_top), Inches(col_w), Inches(header_h))
        h_shape.fill.solid()
        h_shape.fill.fore_color.rgb = BRAND_SLATE
        h_shape.line.fill.background()
        
        # Badge
        badge_dia = 0.26
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(c_left + 0.12), Inches(c_top + 0.09), Inches(badge_dia), Inches(badge_dia))
        badge.fill.solid()
        badge.fill.fore_color.rgb = color
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        pb = tf_b.paragraphs[0]
        pb.alignment = PP_ALIGN.CENTER
        rb = pb.add_run()
        rb.text = f"{num:02d}"
        rb.font.name = FONT_HEADING
        rb.font.size = Pt(8.0)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        # Header Text
        tb_h = slide.shapes.add_textbox(Inches(c_left + 0.44), Inches(c_top + 0.04), Inches(col_w - 0.52), Inches(header_h - 0.08))
        tf_h = tb_h.text_frame
        tf_h.word_wrap = True
        tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0
        p_h = tf_h.paragraphs[0]
        r_t = p_h.add_run()
        r_t.text = title
        r_t.font.name = FONT_HEADING
        r_t.font.size = Pt(9.5)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        p_sub = tf_h.add_paragraph()
        r_sub = p_sub.add_run()
        r_sub.text = f"{actor_dir}  ·  [{tag}]"
        r_sub.font.name = FONT_BODY
        r_sub.font.size = Pt(8.0)
        r_sub.font.color.rgb = RGBColor(0xD6, 0xCE, 0xBF)
        
        # Content Box
        content_top = c_top + header_h + 0.08
        content_h = row_h - header_h - 0.14
        tb_c = slide.shapes.add_textbox(Inches(c_left + 0.16), Inches(content_top), Inches(col_w - 0.32), Inches(content_h))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        
        for idx_b, b in enumerate(bullets):
            p = tf_c.paragraphs[0] if idx_b == 0 else tf_c.add_paragraph()
            p.space_after = Pt(2.5)
            p.line_spacing = 1.15
            if "：" in b:
                parts = b.split("：", 1)
                rh = p.add_run()
                rh.text = "• " + parts[0] + "："
                rh.font.name = FONT_HEADING
                rh.font.size = Pt(8.5)
                rh.font.bold = True
                rh.font.color.rgb = BRAND_SLATE
                
                rb = p.add_run()
                rb.text = parts[1]
                rb.font.name = FONT_BODY
                rb.font.size = Pt(8.0)
                rb.font.color.rgb = TEXT_BODY
            else:
                rb = p.add_run()
                rb.text = "• " + b
                rb.font.name = FONT_BODY
                rb.font.size = Pt(8.0)
                rb.font.color.rgb = TEXT_BODY

print("Helper functions defined successfully!")
