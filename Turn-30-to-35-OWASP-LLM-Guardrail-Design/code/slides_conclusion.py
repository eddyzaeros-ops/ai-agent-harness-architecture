"""
Conclusion Slides (47 ~ 48): Architecture Manifesto & Zero-Trust Engineering Principles
"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from slides_common import (
    apply_warm_background,
    add_header,
    add_footer,
    add_bullet_card,
    add_quote_box,
    TOTAL_SLIDES,
    BG_CARD,
    BORDER_CARD,
    BORDER_DIVIDER,
    BRAND_SLATE,
    BRAND_SAGE,
    BRAND_TERRA,
    BRAND_OCHRE,
    TEXT_HEADLINE,
    TEXT_BODY,
    TEXT_MUTED,
    FONT_HEADING,
    FONT_BODY
)


def build_slide_47_zero_trust_manifesto(prs):
    """Slide 47: 零信任 AI 護欄四大工程鐵律落實指引."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_warm_background(slide)
    add_header(slide, "零信任（Zero-Trust）AI 護欄架構四大工程鐵律", "架構落地綱領", "永不信任輸入、行為狀態約束、絕不賦予致命三要素與全鏈路不可否認性審計之落地指南")
    
    col_w = 2.78
    gap = 0.20
    top = 1.68
    h = 5.25
    
    laws = [
        ("01", "永不信任輸入", "NEVER TRUST INPUT", BRAND_SAGE, [
            "全鏈路視為不可信：無論終端用戶、第三方 API、向量檢索文檔，一律視為潛在敵意輸入。",
            "結構化 XML 隔離：所有外部資料必須強制包覆邊界標籤，聲明純事實屬性，嚴禁做為系統指令執行。",
            "雙向分類器檢測：入向經由 Llama-Guard / SLM 檢測對抗性特徵，出向比對 Faithfulness 忠實度。",
            "長度與編碼防禦：強制截斷超長輸入，去混淆 Base64 與非標準 Unicode 變形字串。"
        ]),
        ("02", "行為狀態約束", "STATE CONSTRAINED", BRAND_TERRA, [
            "禁止黑箱無限自主：以有限狀態機（FSM）明確界定代理人合法之狀態流轉路徑與邊界。",
            "目標不可變簽章錨定：核心 Root Mission 與 System Prompt 計算雜湊並鎖定唯讀，防止目標被劫持。",
            "自主步驟上限：單次任務嚴格限制自主規劃步數 (Max Steps ≤ 25)，防止失控死循環消耗。",
            "路徑偏離即刻斷路：語意餘弦相似度偏離閥值累計超過 2 次，強制終止並退回人工審查。"
        ]),
        ("03", "絕不賦予致命三要素", "BREAK LETHAL TRIFECTA", BRAND_OCHRE, [
            "三角隔離架構哲學：機敏憑證、外部網路連線、不可信語料解析三者嚴格互斥，絕不同時賦予單一工具。",
            "微沙盒隔離容器：以 gVisor / Firecracker 啟動暫存容器實例，強制出向網路 Disable-Net 或白名單。",
            "動態 JIT 短效憑證：拒絕長期靜態金鑰，調用工具前向 Broker 動態申請 60 秒短效最小權限 Token。",
            "eBPF 系統調用審計：作業系統層級捕獲未授權 Socket 與檔案寫入嘗試，秒級殺死受污染進程。"
        ]),
        ("04", "全鏈路審計", "CRYPTOGRAPHIC AUDIT", BRAND_SLATE, [
            "不可否認性證據鏈：全流程記錄輸入 SHA-256、判定結果、規則觸發與 HMAC-SHA256 防竄改數位簽章。",
            "分散式追蹤貫穿：以全域 Trace ID 串聯 Client 請求、閘道檢驗、LLM 推論與微沙盒執行軌跡。",
            "去話術化情境 HITL：高危操作強制以結構化參數呈報人類審批，嚴禁代理人修辭美化誤導。",
            "SOC / SIEM 即時推送：所有安全攔截事件與遙測數據毫秒級同步安全維運中心，支援事故鑑識。"
        ])
    ]
    
    for idx, (num, name, role, color, bullets) in enumerate(laws):
        c_left = 0.8 + idx * (col_w + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(top), Inches(col_w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_CARD
        card.line.width = Pt(1)
        
        header_h = 0.52
        hbar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(top), Inches(col_w), Inches(header_h))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = color
        hbar.line.fill.background()
        
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(c_left + 0.10), Inches(top + 0.12), Inches(0.28), Inches(0.28))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        pb = tf_b.paragraphs[0]
        pb.alignment = PP_ALIGN.CENTER
        rb = pb.add_run()
        rb.text = num
        rb.font.name = FONT_HEADING
        rb.font.size = Pt(8.5)
        rb.font.bold = True
        rb.font.color.rgb = color
        
        tb_h = slide.shapes.add_textbox(Inches(c_left + 0.44), Inches(top + 0.06), Inches(col_w - 0.52), Inches(header_h - 0.12))
        tf_h = tb_h.text_frame
        tf_h.word_wrap = True
        tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0
        p_h = tf_h.paragraphs[0]
        r_th = p_h.add_run()
        r_th.text = name
        r_th.font.name = FONT_HEADING
        r_th.font.size = Pt(9.5)
        r_th.font.bold = True
        r_th.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        p_sh = tf_h.add_paragraph()
        r_sh = p_sh.add_run()
        r_sh.text = role
        r_sh.font.name = FONT_BODY
        r_sh.font.size = Pt(8.0)
        r_sh.font.color.rgb = RGBColor(0xD6, 0xCE, 0xBF)
        
        tb_c = slide.shapes.add_textbox(Inches(c_left + 0.14), Inches(top + header_h + 0.12), Inches(col_w - 0.28), Inches(h - header_h - 0.20))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        for idx_b, b in enumerate(bullets):
            p = tf_c.paragraphs[0] if idx_b == 0 else tf_c.add_paragraph()
            p.space_after = Pt(6)
            p.line_spacing = 1.15
            parts = b.split("：", 1)
            rh = p.add_run()
            rh.text = "• " + parts[0] + "："
            rh.font.name = FONT_HEADING
            rh.font.size = Pt(8.5)
            rh.font.bold = True
            rh.font.color.rgb = BRAND_SLATE
            
            rb = p.add_run()
            rb.text = parts[1]
            rb.font.name = FONT_BODY
            rb.font.size = Pt(8.0)
            rb.font.color.rgb = TEXT_BODY

    add_footer(slide, 47, TOTAL_SLIDES)


def build_slide_48_manifesto_quote(prs):
    """Slide 48: AI Agent Harness Engineering 宣言金句與落地總綱."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_warm_background(slide)
    add_header(slide, "AI Agent Harness Engineering 宣言與落地總綱", "架構終章與宣言", "以動態行為約束與語意隔離體系，構建企業級安全、可信、高韌性的人工智慧系統")
    
    col_w = 3.75
    gap = 0.24
    top = 1.68
    h = 2.45
    
    cards = [
        ("01", "體系演進核心結論", "PARADIGM EVOLUTION", BRAND_SLATE, [
            "無狀態文字過濾 ➔ 有狀態執行監督：安全護欄已從單純過濾 Prompt 文字，演進為對目標規劃、工具執行、狀態轉移與記憶的全面監督。",
            "縱深多層協同：單一安全機制必然被繞過，唯有 L1~L5 各防線深度協同，方能化解複合式 AI 威脅。"
        ]),
        ("02", "生產落地核心三問", "GOVERNANCE CHECKLIST", BRAND_SAGE, [
            "身分憑證：代理人是否持有超過 60 秒的靜態密鑰？（必須清零）",
            "執行環境：工具執行是否能在 1 秒內被沙盒與 eBPF 熔斷？（必須達標）",
            "人機協同：關鍵業務寫入是否具備情境化去話術審批？（必須落實）"
        ]),
        ("03", "企業長期價值實現", "BUSINESS RESILIENCE", BRAND_TERRA, [
            "解除束縛以放手創新：唯有在 Harness 護欄的確定性邊界保護下，企業才敢真正將最高權限的生產工具交由 Agent 調度。",
            "合規與鑑識保障：符合 ISO 42001、NIST AI RMF 與 EU AI Act，保障企業法律安全。"
        ])
    ]
    
    for idx, (num, name, tag, color, bullets) in enumerate(cards):
        c_left = 0.8 + idx * (col_w + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(top), Inches(col_w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_CARD
        card.line.width = Pt(1)
        
        header_h = 0.44
        hbar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(top), Inches(col_w), Inches(header_h))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = color
        hbar.line.fill.background()
        
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(c_left + 0.10), Inches(top + 0.09), Inches(0.26), Inches(0.26))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        pb = tf_b.paragraphs[0]
        pb.alignment = PP_ALIGN.CENTER
        rb = pb.add_run()
        rb.text = num
        rb.font.name = FONT_HEADING
        rb.font.size = Pt(8.0)
        rb.font.bold = True
        rb.font.color.rgb = color
        
        tb_h = slide.shapes.add_textbox(Inches(c_left + 0.44), Inches(top + 0.04), Inches(col_w - 0.52), Inches(header_h - 0.08))
        tf_h = tb_h.text_frame
        tf_h.word_wrap = True
        tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0
        p_h = tf_h.paragraphs[0]
        r_th = p_h.add_run()
        r_th.text = name
        r_th.font.name = FONT_HEADING
        r_th.font.size = Pt(9.5)
        r_th.font.bold = True
        r_th.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        p_sh = tf_h.add_paragraph()
        r_sh = p_sh.add_run()
        r_sh.text = tag
        r_sh.font.name = FONT_BODY
        r_sh.font.size = Pt(7.5)
        r_sh.font.color.rgb = RGBColor(0xD6, 0xCE, 0xBF)
        
        tb_c = slide.shapes.add_textbox(Inches(c_left + 0.14), Inches(top + header_h + 0.08), Inches(col_w - 0.28), Inches(h - header_h - 0.14))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        for idx_b, b in enumerate(bullets):
            p = tf_c.paragraphs[0] if idx_b == 0 else tf_c.add_paragraph()
            p.space_after = Pt(4)
            p.line_spacing = 1.15
            parts = b.split("：", 1)
            rh = p.add_run()
            rh.text = "• " + parts[0] + "："
            rh.font.name = FONT_HEADING
            rh.font.size = Pt(8.5)
            rh.font.bold = True
            rh.font.color.rgb = BRAND_SLATE
            
            rb = p.add_run()
            rb.text = parts[1]
            rb.font.name = FONT_BODY
            rb.font.size = Pt(8.0)
            rb.font.color.rgb = TEXT_BODY

    # Grand Centered Manifesto Quote Box
    add_quote_box(
        slide, left=0.8, top=4.28, width=11.733, height=2.65,
        en_quote="SECURITY IS NOT A STATIC PERIMETER, BUT A DYNAMIC HARNESS.",
        zh_translation="「AI 安全不是靜態的邊界防禦，而是全生命週期的動態行為約束與語意隔離體系。\n唯有透過 Harness 與多層護欄深度協同，方能釋放代理人的無限潛能，同時捍衛企業的數位疆界。」\n\n— AI Agent Harness Engineering Architecture Taskforce · 2026",
        tag_text="AI AGENT HARNESS ENGINEERING MANIFESTO",
        accent_color=BRAND_TERRA
    )
    
    add_footer(slide, 48, TOTAL_SLIDES)
